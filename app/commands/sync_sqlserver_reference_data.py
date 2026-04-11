"""Manual SQL Server reference-data import command for multiple target databases."""

from __future__ import annotations

import argparse
import csv
import json
import logging
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Sequence

from openpyxl import Workbook, load_workbook
from pydantic import BaseModel, ConfigDict, model_validator

from app.common.services.sqlserver_reference_sync import (
    DEFAULT_SQL_SERVER_PORT,
    SqlServerConnectionSettings,
    SqlServerEngineFactory,
    SqlServerMetadataLoader,
    SqlServerModelSyncExecutor,
    SqlServerSyncPlanBuilder,
    SqlServerSyncSqlBuilder,
    SqlServerSyncTableOverride,
    SqlServerTableReference,
)


@dataclass(frozen=True)
class TargetDatabaseTarget:
    host: str
    port: int
    username: str
    password: str
    source_row: int


@dataclass(frozen=True)
class DataSyncPlan:
    sync_flag: str
    model_name: str
    table: SqlServerTableReference
    data_file_path: Path
    override: SqlServerSyncTableOverride | None = None


@dataclass(frozen=True)
class DataSyncResult:
    target_host: str
    target_database: str
    sync_flag: str
    model_name: str
    data_file: str
    inserted_count: int
    updated_count: int
    skipped_row_count: int
    source_row_count: int | None
    status: str


@dataclass(frozen=True)
class DataSyncError:
    target_host: str | None
    target_database: str | None
    sync_flag: str | None
    model_name: str | None
    source_row: int | None
    error_message: str


@dataclass(frozen=True)
class DataSyncSkippedRow:
    target_host: str
    target_database: str
    sync_flag: str
    model_name: str
    data_file: str
    source_row: int | None
    row_identifier: str | None
    reference_table: str
    reference_column: str
    reference_value: str
    reason: str


@dataclass(frozen=True)
class DataSyncSummary:
    targets_processed: int
    selected_tables: int
    table_runs_attempted: int
    successful_table_runs: int
    skipped_row_count: int
    error_count: int
    output_path: Path


@dataclass(frozen=True)
class MissingReferenceSkipRule:
    source_column_name: str
    reference_table: SqlServerTableReference
    reference_column_name: str = "id"


class SyncDbPayload(BaseModel):
    model_config = ConfigDict(extra="forbid")

    sync_transport_tariffs: bool = False
    sync_transport_year_amounts: bool = False

    @model_validator(mode="after")
    def apply_transport_dependencies(self) -> "SyncDbPayload":
        if self.sync_transport_year_amounts:
            self.sync_transport_tariffs = True
        return self


class TargetServerInputReader:
    HEADER_ALIASES = {
        "host": {"ip", "host", "hostname", "server", "serverip", "serverhost", "address"},
        "port": {"port", "serverport", "dbport"},
        "username": {"username", "user", "userid", "uid", "login"},
        "password": {"password", "pass", "passwd", "pwd"},
    }

    def read(
        self,
        *,
        input_path: Path,
        default_target_database: str,
    ) -> tuple[list[TargetDatabaseTarget], list[DataSyncError]]:
        del default_target_database
        rows = self._read_rows(input_path=input_path)
        targets: list[TargetDatabaseTarget] = []
        errors: list[DataSyncError] = []
        for row_number, row in rows:
            try:
                targets.append(
                    TargetDatabaseTarget(
                        host=self._require_value(row=row, key="host", field_label="host"),
                        port=self._safe_port(row.get("port")) or DEFAULT_SQL_SERVER_PORT,
                        username=self._require_value(
                            row=row,
                            key="username",
                            field_label="username",
                        ),
                        password=self._require_value(
                            row=row,
                            key="password",
                            field_label="password",
                        ),
                        source_row=row_number,
                    )
                )
            except ValueError as exc:
                errors.append(
                    DataSyncError(
                        target_host=self._normalize_text(row.get("host")),
                        target_database=None,
                        sync_flag=None,
                        model_name=None,
                        source_row=row_number,
                        error_message=f"Input row {row_number}: {exc}",
                    )
                )
        return targets, errors

    def _read_rows(self, *, input_path: Path) -> list[tuple[int, dict[str, Any]]]:
        if input_path.suffix.lower() == ".csv":
            with input_path.open("r", encoding="utf-8-sig", newline="") as file_handle:
                rows = list(csv.reader(file_handle))
        elif input_path.suffix.lower() == ".xlsx":
            workbook = load_workbook(filename=input_path, read_only=True, data_only=True)
            try:
                rows = list(workbook.active.iter_rows(values_only=True))
            finally:
                workbook.close()
        else:
            raise ValueError("Target input file must be .xlsx or .csv.")
        if not rows:
            return []
        header_lookup = self._build_header_lookup(headers=rows[0])
        return [
            (row_number, self._extract_row_values(row=row, header_lookup=header_lookup))
            for row_number, row in enumerate(rows[1:], start=2)
        ]

    def _build_header_lookup(self, *, headers: Sequence[object]) -> dict[str, int]:
        lookup: dict[str, int] = {}
        for index, header in enumerate(headers):
            normalized_header = self._normalize_header_name(header)
            if normalized_header is None:
                continue
            for canonical_name, aliases in self.HEADER_ALIASES.items():
                if normalized_header in aliases and canonical_name not in lookup:
                    lookup[canonical_name] = index
        return lookup

    def _extract_row_values(
        self,
        *,
        row: Sequence[object],
        header_lookup: dict[str, int],
    ) -> dict[str, Any]:
        return {
            canonical_name: row[column_index] if column_index < len(row) else None
            for canonical_name, column_index in header_lookup.items()
        }

    def _require_value(self, *, row: dict[str, Any], key: str, field_label: str) -> str:
        normalized_value = self._normalize_text(row.get(key))
        if normalized_value is None:
            raise ValueError(f"Missing {field_label} value.")
        return normalized_value

    def _normalize_header_name(self, value: object) -> str | None:
        text_value = self._normalize_text(value)
        if text_value is None:
            return None
        return "".join(character for character in text_value.lower() if character.isalnum())

    def _normalize_text(self, value: object | None) -> str | None:
        if value is None:
            return None
        text_value = str(value).strip()
        if not text_value:
            return None
        return text_value

    def _safe_port(self, value: object | None) -> int | None:
        if value is None:
            return None
        try:
            return int(str(value).strip())
        except ValueError:
            return None


class ReferenceDataFileReader:
    def read(self, input_path: Path) -> list[dict[str, Any]]:
        rows = self._read_raw_rows(input_path=input_path)
        if not rows:
            return []
        headers = [self._normalize_header_name(value) for value in rows[0]]
        column_names = [header for header in headers if header]
        normalized_rows: list[dict[str, Any]] = []
        for row_number, row in enumerate(rows[1:], start=2):
            payload: dict[str, Any] = {}
            for index, column_name in enumerate(headers):
                if not column_name:
                    continue
                payload[column_name] = row[index] if index < len(row) else None
            payload["__source_row__"] = row_number
            normalized_rows.append(payload)
        if not column_names:
            raise ValueError(f"Data file '{input_path}' does not contain usable headers.")
        return normalized_rows

    def _read_raw_rows(self, *, input_path: Path) -> list[Sequence[object]]:
        if input_path.suffix.lower() == ".csv":
            with input_path.open("r", encoding="utf-8-sig", newline="") as file_handle:
                return list(csv.reader(file_handle))
        if input_path.suffix.lower() == ".xlsx":
            workbook = load_workbook(filename=input_path, read_only=True, data_only=True)
            try:
                return list(workbook.active.iter_rows(values_only=True))
            finally:
                workbook.close()
        raise ValueError("Data file must be .xlsx or .csv.")

    def _normalize_header_name(self, value: object | None) -> str | None:
        if value is None:
            return None
        text_value = str(value).strip()
        if not text_value:
            return None
        return text_value


class SyncPlanBuilder:
    _TABLE_ARGS = {
        "TransportTariffs": "transport_tariffs_file",
        "TransportTariffLocals": "transport_tariff_locals_file",
        "TransportYearAmounts": "transport_year_amounts_file",
        "TransportYearAmountLocals": "transport_year_amount_locals_file",
    }
    _TABLE_SCHEMAS = {
        "TransportTariffs": "Trp",
        "TransportTariffLocals": "Trp",
        "TransportYearAmounts": "Trp",
        "TransportYearAmountLocals": "Trp",
    }
    _TABLE_OVERRIDES = {
        "TransportYearAmounts": SqlServerSyncTableOverride(
            schema_name="Trp",
            column_names=("Amount", "Year", "IsDeleted", "TransportTariffId"),
            primary_key_columns=("Year", "TransportTariffId"),
        ),
    }
    _GROUPS: tuple[tuple[str, tuple[str, ...]], ...] = (
        ("sync_transport_tariffs", ("TransportTariffs", "TransportTariffLocals")),
        ("sync_transport_year_amounts", ("TransportYearAmounts", "TransportYearAmountLocals")),
    )

    def build(self, *, payload: SyncDbPayload, args: argparse.Namespace) -> list[DataSyncPlan]:
        plans: list[DataSyncPlan] = []
        payload_data = payload.model_dump()
        for sync_flag, table_names in self._GROUPS:
            if not payload_data[sync_flag]:
                continue
            for table_name in table_names:
                file_path = getattr(args, self._TABLE_ARGS[table_name])
                if not file_path:
                    continue
                plans.append(
                    DataSyncPlan(
                        sync_flag=sync_flag,
                        model_name=table_name,
                        table=SqlServerTableReference(
                            model_name=table_name,
                            schema_name=self._TABLE_SCHEMAS[table_name],
                            table_name=table_name,
                        ),
                        data_file_path=Path(file_path).resolve(),
                        override=self._TABLE_OVERRIDES.get(table_name),
                    )
                )
        return plans


class SyncReportWriter:
    RESULT_COLUMNS = [
        "target_host",
        "target_database",
        "sync_flag",
        "model_name",
        "data_file",
        "inserted_count",
        "updated_count",
        "skipped_row_count",
        "source_row_count",
        "status",
    ]
    ERROR_COLUMNS = [
        "target_host",
        "target_database",
        "sync_flag",
        "model_name",
        "source_row",
        "error_message",
    ]
    SKIPPED_COLUMNS = [
        "target_host",
        "target_database",
        "sync_flag",
        "model_name",
        "data_file",
        "source_row",
        "row_identifier",
        "reference_table",
        "reference_column",
        "reference_value",
        "reason",
    ]

    def save(
        self,
        *,
        output_path: Path,
        summary: DataSyncSummary,
        results: Sequence[DataSyncResult],
        errors: Sequence[DataSyncError],
        skipped_rows: Sequence[DataSyncSkippedRow],
    ) -> None:
        output_path.parent.mkdir(parents=True, exist_ok=True)
        if output_path.suffix.lower() == ".json":
            output_path.write_text(
                json.dumps(
                    {
                        "summary": summary.__dict__ | {"output_path": str(summary.output_path)},
                        "results": [result.__dict__ for result in results],
                        "errors": [error.__dict__ for error in errors],
                        "skipped_rows": [skipped_row.__dict__ for skipped_row in skipped_rows],
                    },
                    ensure_ascii=False,
                    indent=2,
                    default=str,
                ),
                encoding="utf-8",
            )
            return

        workbook = Workbook(write_only=True)
        summary_sheet = workbook.create_sheet(title="summary")
        results_sheet = workbook.create_sheet(title="results")
        errors_sheet = workbook.create_sheet(title="errors")
        skipped_sheet = workbook.create_sheet(title="skipped_rows")
        summary_sheet.append(["metric", "value"])
        for key, value in summary.__dict__.items():
            summary_sheet.append([key, str(value)])
        results_sheet.append(self.RESULT_COLUMNS)
        for result in results:
            results_sheet.append([getattr(result, column_name) for column_name in self.RESULT_COLUMNS])
        errors_sheet.append(self.ERROR_COLUMNS)
        for error in errors:
            errors_sheet.append([getattr(error, column_name) for column_name in self.ERROR_COLUMNS])
        skipped_sheet.append(self.SKIPPED_COLUMNS)
        for skipped_row in skipped_rows:
            skipped_sheet.append(
                [getattr(skipped_row, column_name) for column_name in self.SKIPPED_COLUMNS]
            )
        workbook.save(output_path)


class SyncSqlServerReferenceDataCommand:
    _MISSING_REFERENCE_SKIP_RULES = {
        "TransportYearAmounts": MissingReferenceSkipRule(
            source_column_name="TransportTariffId",
            reference_table=SqlServerTableReference(
                model_name="TransportTariffs",
                schema_name="Trp",
                table_name="TransportTariffs",
            ),
        ),
        "TransportYearAmountLocals": MissingReferenceSkipRule(
            source_column_name="TransportTariffLocalId",
            reference_table=SqlServerTableReference(
                model_name="TransportTariffLocals",
                schema_name="Trp",
                table_name="TransportTariffLocals",
            ),
        ),
    }

    def __init__(
        self,
        *,
        target_reader: TargetServerInputReader,
        data_reader: ReferenceDataFileReader,
        plan_builder: SyncPlanBuilder,
        executor: SqlServerModelSyncExecutor,
        report_writer: SyncReportWriter,
        logger: logging.Logger,
    ) -> None:
        self._target_reader = target_reader
        self._data_reader = data_reader
        self._plan_builder = plan_builder
        self._executor = executor
        self._report_writer = report_writer
        self._logger = logger
        self._last_errors: list[DataSyncError] = []
        self._last_skipped_rows: list[DataSyncSkippedRow] = []

    def run(
        self,
        *,
        targets_input_path: Path,
        output_path: Path,
        payload: SyncDbPayload,
        args: argparse.Namespace,
    ) -> DataSyncSummary:
        targets, errors = self._target_reader.read(
            input_path=targets_input_path,
            default_target_database="",
        )
        plans = self._plan_builder.build(payload=payload, args=args)
        data_by_file = {
            plan.data_file_path: self._data_reader.read(plan.data_file_path)
            for plan in plans
        }
        results: list[DataSyncResult] = []
        skipped_rows: list[DataSyncSkippedRow] = []
        successful_table_runs = 0
        table_runs_attempted = 0

        for target in targets:
            discovery_settings = SqlServerConnectionSettings(
                host=target.host,
                port=target.port,
                username=target.username,
                password=target.password,
                database_name="master",
            )
            try:
                self._executor.check_target_connection(target_settings=discovery_settings)
                database_names = self._executor.discover_target_databases(
                    target_settings=discovery_settings
                )
            except Exception as exc:
                self._logger.error(
                    "target_connection_failed host=%s port=%s source_row=%s error=%s",
                    target.host,
                    target.port,
                    target.source_row,
                    exc,
                )
                errors.append(
                    DataSyncError(
                        target_host=target.host,
                        target_database=None,
                        sync_flag=None,
                        model_name=None,
                        source_row=target.source_row,
                        error_message=str(exc),
                    )
                )
                continue

            for database_name in database_names:
                target_settings = SqlServerConnectionSettings(
                    host=target.host,
                    port=target.port,
                    username=target.username,
                    password=target.password,
                    database_name=database_name,
                )
                for plan in plans:
                    table_runs_attempted += 1
                    try:
                        filtered_row_payloads = data_by_file[plan.data_file_path]
                        run_skipped_rows: list[DataSyncSkippedRow] = []
                        if plan.model_name in self._MISSING_REFERENCE_SKIP_RULES:
                            filtered_row_payloads, run_skipped_rows = self._split_skipped_rows(
                                target_settings=target_settings,
                                plan=plan,
                                row_payloads=filtered_row_payloads,
                            )
                            skipped_rows.extend(run_skipped_rows)
                        stats = self._executor.execute_from_rows(
                            target_settings=target_settings,
                            table=plan.table,
                            row_payloads=filtered_row_payloads,
                            override=plan.override,
                        )
                        successful_table_runs += 1
                        results.append(
                            DataSyncResult(
                                target_host=target.host,
                                target_database=database_name,
                                sync_flag=plan.sync_flag,
                                model_name=plan.model_name,
                                data_file=str(plan.data_file_path),
                                inserted_count=stats.inserted_count,
                                updated_count=stats.updated_count,
                                skipped_row_count=len(run_skipped_rows),
                                source_row_count=stats.source_row_count,
                                status="ok",
                            )
                        )
                    except Exception as exc:
                        self._logger.error(
                            "table_sync_failed host=%s database=%s table=%s source_row=%s error=%s",
                            target.host,
                            database_name,
                            plan.model_name,
                            target.source_row,
                            exc,
                        )
                        errors.append(
                            DataSyncError(
                                target_host=target.host,
                                target_database=database_name,
                                sync_flag=plan.sync_flag,
                                model_name=plan.model_name,
                                source_row=target.source_row,
                                error_message=str(exc),
                            )
                        )

        summary = DataSyncSummary(
            targets_processed=len(targets),
            selected_tables=len(plans),
            table_runs_attempted=table_runs_attempted,
            successful_table_runs=successful_table_runs,
            skipped_row_count=len(skipped_rows),
            error_count=len(errors),
            output_path=output_path,
        )
        self._report_writer.save(
            output_path=output_path,
            summary=summary,
            results=results,
            errors=errors,
            skipped_rows=skipped_rows,
        )
        self._last_errors = list(errors)
        self._last_skipped_rows = list(skipped_rows)
        return summary

    @property
    def last_errors(self) -> tuple[DataSyncError, ...]:
        return tuple(self._last_errors)

    @property
    def last_skipped_rows(self) -> tuple[DataSyncSkippedRow, ...]:
        return tuple(self._last_skipped_rows)

    def _split_skipped_rows(
        self,
        *,
        target_settings: SqlServerConnectionSettings,
        plan: DataSyncPlan,
        row_payloads: Sequence[dict[str, Any]],
    ) -> tuple[list[dict[str, Any]], list[DataSyncSkippedRow]]:
        skip_rule = self._MISSING_REFERENCE_SKIP_RULES[plan.model_name]
        existing_reference_values = {
            self._normalize_lookup_value(value)
            for value in self._executor.fetch_existing_column_values(
                target_settings=target_settings,
                table=skip_rule.reference_table,
                column_name=skip_rule.reference_column_name,
            )
        }
        valid_rows: list[dict[str, Any]] = []
        skipped_rows: list[DataSyncSkippedRow] = []
        for row_payload in row_payloads:
            reference_value = row_payload.get(skip_rule.source_column_name)
            normalized_reference_value = self._normalize_lookup_value(reference_value)
            if normalized_reference_value is None or normalized_reference_value in existing_reference_values:
                valid_rows.append(row_payload)
                continue
            skipped_rows.append(
                DataSyncSkippedRow(
                    target_host=target_settings.host,
                    target_database=target_settings.database_name,
                    sync_flag=plan.sync_flag,
                    model_name=plan.model_name,
                    data_file=str(plan.data_file_path),
                    source_row=self._read_source_row(row_payload),
                    row_identifier=self._build_row_identifier(
                        plan=plan,
                        row_payload=row_payload,
                    ),
                    reference_table=(
                        f"{skip_rule.reference_table.schema_name}."
                        f"{skip_rule.reference_table.table_name}"
                    ),
                    reference_column=skip_rule.source_column_name,
                    reference_value=normalized_reference_value,
                    reason=(
                        f"{skip_rule.source_column_name}={normalized_reference_value} "
                        f"در جدول {skip_rule.reference_table.schema_name}."
                        f"{skip_rule.reference_table.table_name} پیدا نشد و این ردیف skip شد."
                    ),
                )
            )
        return valid_rows, skipped_rows

    def _read_source_row(self, row_payload: dict[str, Any]) -> int | None:
        source_row = row_payload.get("__source_row__")
        if isinstance(source_row, int):
            return source_row
        return None

    def _build_row_identifier(
        self,
        *,
        plan: DataSyncPlan,
        row_payload: dict[str, Any],
    ) -> str | None:
        if "id" in row_payload and row_payload.get("id") is not None:
            normalized_id = self._normalize_lookup_value(row_payload.get("id"))
            if normalized_id is not None:
                return f"id={normalized_id}"
        if plan.model_name == "TransportYearAmounts":
            year = self._normalize_lookup_value(row_payload.get("Year"))
            tariff_id = self._normalize_lookup_value(row_payload.get("TransportTariffId"))
            if year is not None and tariff_id is not None:
                return f"Year={year}, TransportTariffId={tariff_id}"
        return None

    def _normalize_lookup_value(self, value: object | None) -> str | None:
        if value is None:
            return None
        if isinstance(value, bool):
            return "1" if value else "0"
        if isinstance(value, int):
            return str(value)
        if isinstance(value, float):
            if value.is_integer():
                return str(int(value))
            return format(value, "f").rstrip("0").rstrip(".")
        text_value = str(value).strip()
        if not text_value:
            return None
        try:
            decimal_value = Decimal(text_value)
        except InvalidOperation:
            return text_value
        normalized_decimal = decimal_value.normalize()
        if normalized_decimal == normalized_decimal.to_integral():
            return str(normalized_decimal.quantize(Decimal("1")))
        return format(normalized_decimal, "f").rstrip("0").rstrip(".")


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Import reference-data files into multiple SQL Server target databases.",
    )
    parser.add_argument("--targets-input", required=True, help="Path to target servers CSV/XLSX.")
    parser.add_argument("--output", required=True, help="Path to the output report (.xlsx or .json).")
    parser.add_argument("--driver", default="ODBC Driver 18 for SQL Server", help="ODBC driver name.")
    parser.add_argument("--connect-timeout", type=int, default=15, help="Connection timeout in seconds.")
    parser.add_argument("--batch-size", type=int, default=1000, help="Batch size.")
    parser.add_argument("--log-level", default="INFO", choices=["DEBUG", "INFO", "WARNING", "ERROR"])
    parser.add_argument("--sync-transport-tariffs", action="store_true")
    parser.add_argument("--sync-transport-year-amounts", action="store_true")
    parser.add_argument("--transport-tariffs-file", help="Path to TransportTariffs data file.")
    parser.add_argument("--transport-tariff-locals-file", help="Path to TransportTariffLocals data file.")
    parser.add_argument("--transport-year-amounts-file", help="Path to TransportYearAmounts data file.")
    parser.add_argument("--transport-year-amount-locals-file", help="Path to TransportYearAmountLocals data file.")
    return parser


def build_payload_from_args(args: argparse.Namespace) -> SyncDbPayload:
    return SyncDbPayload(
        sync_transport_tariffs=bool(args.sync_transport_tariffs),
        sync_transport_year_amounts=bool(args.sync_transport_year_amounts),
    )


def configure_logging(*, log_level: str) -> logging.Logger:
    logging.basicConfig(
        level=getattr(logging, log_level.upper()),
        format="%(asctime)s %(levelname)s %(message)s",
    )
    return logging.getLogger("sqlserver_reference_import")


def format_summary(summary: DataSyncSummary) -> str:
    return (
        f"targets_processed={summary.targets_processed}\n"
        f"selected_tables={summary.selected_tables}\n"
        f"table_runs_attempted={summary.table_runs_attempted}\n"
        f"successful_table_runs={summary.successful_table_runs}\n"
        f"skipped_row_count={summary.skipped_row_count}\n"
        f"error_count={summary.error_count}\n"
        f"output_file={summary.output_path}"
    )


def format_errors(errors: Sequence[DataSyncError], *, max_items: int = 10) -> str:
    if not errors:
        return ""
    lines = ["errors:"]
    for index, error in enumerate(errors[:max_items], start=1):
        lines.append(
            f"{index}. target_host={error.target_host} "
            f"target_database={error.target_database} "
            f"sync_flag={error.sync_flag} "
            f"model_name={error.model_name} "
            f"source_row={error.source_row} "
            f"error_message={error.error_message}"
        )
    if len(errors) > max_items:
        lines.append(f"... and {len(errors) - max_items} more error(s)")
    return "\n".join(lines)


def format_skipped_rows(
    skipped_rows: Sequence[DataSyncSkippedRow],
    *,
    max_items: int = 10,
) -> str:
    if not skipped_rows:
        return ""
    lines = ["skipped_rows:"]
    for index, skipped_row in enumerate(skipped_rows[:max_items], start=1):
        lines.append(
            f"{index}. target_host={skipped_row.target_host} "
            f"target_database={skipped_row.target_database} "
            f"sync_flag={skipped_row.sync_flag} "
            f"model_name={skipped_row.model_name} "
            f"source_row={skipped_row.source_row} "
            f"row_identifier={skipped_row.row_identifier} "
            f"reference_value={skipped_row.reference_value} "
            f"reason={skipped_row.reason}"
        )
    if len(skipped_rows) > max_items:
        lines.append(f"... and {len(skipped_rows) - max_items} more skipped row(s)")
    return "\n".join(lines)


def format_command_output(
    summary: DataSyncSummary,
    errors: Sequence[DataSyncError],
    skipped_rows: Sequence[DataSyncSkippedRow],
) -> str:
    output_sections = [format_summary(summary)]
    skipped_rows_text = format_skipped_rows(skipped_rows)
    if skipped_rows_text:
        output_sections.append(skipped_rows_text)
    errors_text = format_errors(errors)
    if errors_text:
        output_sections.append(errors_text)
    return "\n".join(output_sections)


def main(argv: Sequence[str] | None = None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)
    payload = build_payload_from_args(args)
    logger = configure_logging(log_level=args.log_level)
    command = SyncSqlServerReferenceDataCommand(
        target_reader=TargetServerInputReader(),
        data_reader=ReferenceDataFileReader(),
        plan_builder=SyncPlanBuilder(),
        executor=SqlServerModelSyncExecutor(
            engine_factory=SqlServerEngineFactory(
                driver=args.driver,
                connect_timeout=args.connect_timeout,
            ),
            metadata_loader=SqlServerMetadataLoader(),
            sync_plan_builder=SqlServerSyncPlanBuilder(),
            sql_builder=SqlServerSyncSqlBuilder(),
            batch_size=args.batch_size,
        ),
        report_writer=SyncReportWriter(),
        logger=logger,
    )
    summary = command.run(
        targets_input_path=Path(args.targets_input).resolve(),
        output_path=Path(args.output).resolve(),
        payload=payload,
        args=args,
    )
    print(format_command_output(summary, command.last_errors, command.last_skipped_rows))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
