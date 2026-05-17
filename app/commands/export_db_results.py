"""Manual SQL Server export command for scanning multiple servers and databases."""

from __future__ import annotations

import argparse
import csv
import json
import logging
import re
import tempfile
from contextlib import contextmanager
from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
from typing import Any, Iterator, Sequence
from uuid import UUID

from openpyxl import Workbook, load_workbook
from sqlalchemy import URL, create_engine, text

DEFAULT_SQL_SERVER_PORT = 1433
SYSTEM_DATABASES = {"master", "model", "msdb", "tempdb"}
READ_ONLY_DISALLOWED_TOKENS = {
    "alter",
    "backup",
    "begin",
    "commit",
    "create",
    "dbcc",
    "delete",
    "detach",
    "drop",
    "exec",
    "execute",
    "grant",
    "insert",
    "into",
    "kill",
    "merge",
    "restore",
    "rollback",
    "revoke",
    "shutdown",
    "truncate",
    "update",
    "use",
}
READ_ONLY_ALLOWED_PREFIXES = {"select", "with", "declare", "set"}
SQL_BLOCK_COMMENT_PATTERN = re.compile(r"/\*.*?\*/", re.DOTALL)
SQL_LINE_COMMENT_PATTERN = re.compile(r"--[^\r\n]*")
SQL_STRING_LITERAL_PATTERN = re.compile(r"N?'(?:''|[^'])*'")
SQL_TOKEN_PATTERN = re.compile(r"[A-Za-z_][A-Za-z0-9_]*")
SCAN_STATUS_QUERY_OK = "ok"
SCAN_STATUS_NO_ROWS = "no_rows"
SCAN_STATUS_WRITE_OK = "write_ok"


@dataclass(frozen=True)
class ServerTarget:
    host: str
    port: int
    username: str
    password: str
    source_row: int


@dataclass(frozen=True)
class ExportError:
    server_ip: str | None
    port: int | None
    database_name: str | None
    error_message: str


@dataclass(frozen=True)
class ExportSummary:
    servers_processed: int
    databases_scanned: int
    success_count: int
    error_count: int
    output_path: Path


@dataclass(frozen=True)
class QueryExecutionPayload:
    columns: list[str]
    rows: Iterator[Sequence[Any]] | None
    affected_row_count: int | None


class ServerInputReader:
    HEADER_ALIASES = {
        "ip": {"ip", "host", "hostname", "server", "serverip", "serverhost", "address"},
        "port": {"port", "serverport", "dbport"},
        "username": {"username", "user", "userid", "uid", "login"},
        "password": {"password", "pass", "passwd", "pwd"},
    }

    def read(self, input_path: Path) -> tuple[list[ServerTarget], list[ExportError]]:
        suffix = input_path.suffix.lower()
        if suffix == ".csv":
            rows = self._read_csv(input_path=input_path)
        elif suffix == ".xlsx":
            rows = self._read_xlsx(input_path=input_path)
        else:
            raise ValueError("Input file must be .xlsx or .csv.")

        servers: list[ServerTarget] = []
        errors: list[ExportError] = []
        for row_number, row in rows:
            try:
                servers.append(self._build_server_target(row=row, row_number=row_number))
            except ValueError as exc:
                errors.append(
                    ExportError(
                        server_ip=self._normalize_text(row.get("ip")),
                        port=self._safe_port(row.get("port")),
                        database_name=None,
                        error_message=f"Input row {row_number}: {exc}",
                    )
                )
        return servers, errors

    def _read_csv(self, *, input_path: Path) -> list[tuple[int, dict[str, Any]]]:
        with input_path.open("r", encoding="utf-8-sig", newline="") as file_handle:
            reader = csv.reader(file_handle)
            rows = list(reader)
        if not rows:
            return []
        header_lookup = self._build_header_lookup(headers=rows[0])
        return [
            (row_number, self._extract_row_values(row=row, header_lookup=header_lookup))
            for row_number, row in enumerate(rows[1:], start=2)
        ]

    def _read_xlsx(self, *, input_path: Path) -> list[tuple[int, dict[str, Any]]]:
        workbook = load_workbook(filename=input_path, read_only=True, data_only=True)
        try:
            sheet = workbook.active
            rows = list(sheet.iter_rows(values_only=True))
        finally:
            workbook.close()
        if not rows:
            return []
        header_lookup = self._build_header_lookup(headers=rows[0])
        return [
            (row_number, self._extract_row_values(row=row, header_lookup=header_lookup))
            for row_number, row in enumerate(rows[1:], start=2)
        ]

    def _build_header_lookup(self, *, headers: Sequence[object]) -> dict[str, int]:
        lookup: dict[str, int] = {}
        for index, header in enumerate(headers):
            normalized_header = self._normalize_header_name(header)
            if normalized_header is None:
                continue
            for canonical_name, aliases in self.HEADER_ALIASES.items():
                if normalized_header in aliases and canonical_name not in lookup:
                    lookup[canonical_name] = index
        return lookup

    def _extract_row_values(
        self,
        *,
        row: Sequence[object],
        header_lookup: dict[str, int],
    ) -> dict[str, Any]:
        return {
            canonical_name: (
                row[column_index]
                if column_index < len(row)
                else None
            )
            for canonical_name, column_index in header_lookup.items()
        }

    def _build_server_target(
        self,
        *,
        row: dict[str, Any],
        row_number: int,
    ) -> ServerTarget:
        host = self._normalize_text(row.get("ip"))
        username = self._normalize_text(row.get("username"))
        password = self._normalize_text(row.get("password"))
        port = self._safe_port(row.get("port")) or DEFAULT_SQL_SERVER_PORT

        if host is None:
            raise ValueError("Missing ip/host value.")
        if username is None:
            raise ValueError("Missing username value.")
        if password is None:
            raise ValueError("Missing password value.")

        return ServerTarget(
            host=host,
            port=port,
            username=username,
            password=password,
            source_row=row_number,
        )

    def _normalize_header_name(self, value: object) -> str | None:
        text_value = self._normalize_text(value)
        if text_value is None:
            return None
        return "".join(character for character in text_value.lower() if character.isalnum())

    def _normalize_text(self, value: object | None) -> str | None:
        if value is None:
            return None
        text_value = str(value).strip()
        if not text_value:
            return None
        return text_value

    def _safe_port(self, value: object | None) -> int | None:
        if value is None:
            return None
        try:
            return int(str(value).strip())
        except ValueError:
            return None


class SqlServerConnector:
    def __init__(
        self,
        *,
        driver: str,
        connect_timeout: int,
    ) -> None:
        self._driver = driver
        self._connect_timeout = connect_timeout

    def discover_databases(self, *, server: ServerTarget) -> list[str]:
        discovery_query = text(
            """
            SELECT name
            FROM sys.databases
            WHERE state = 0
              AND name NOT IN ('master', 'model', 'msdb', 'tempdb')
            ORDER BY name
            """
        )
        engine = self._create_engine(server=server, database_name="master")
        try:
            with engine.connect() as connection:
                return [str(name) for name in connection.execute(discovery_query).scalars().all()]
        finally:
            engine.dispose()

    @contextmanager
    def execute_query(
        self,
        *,
        server: ServerTarget,
        database_name: str,
        query: str,
        allow_write: bool,
    ) -> Iterator[QueryExecutionPayload]:
        engine = self._create_engine(server=server, database_name=database_name)
        try:
            with engine.connect() as connection:
                transaction = connection.begin()
                try:
                    result = connection.execute(text(query))
                    if result.returns_rows:
                        yield QueryExecutionPayload(
                            columns=list(result.keys()),
                            rows=result,
                            affected_row_count=None,
                        )
                    else:
                        affected_row_count = result.rowcount if result.rowcount >= 0 else None
                        yield QueryExecutionPayload(
                            columns=[],
                            rows=None,
                            affected_row_count=affected_row_count,
                        )
                    if allow_write:
                        transaction.commit()
                finally:
                    if transaction.is_active:
                        transaction.rollback()
        finally:
            engine.dispose()

    def _create_engine(self, *, server: ServerTarget, database_name: str):
        self._ensure_pyodbc_available()
        url = URL.create(
            "mssql+pyodbc",
            username=server.username,
            password=server.password,
            host=server.host,
            port=server.port,
            database=database_name,
            query={
                "driver": self._driver,
                "TrustServerCertificate": "yes",
                "Encrypt": "no",
            },
        )
        return create_engine(
            url,
            pool_pre_ping=True,
            hide_parameters=True,
            connect_args={"timeout": self._connect_timeout},
        )

    def _ensure_pyodbc_available(self) -> None:
        try:
            import pyodbc  # noqa: F401
        except ImportError as exc:
            raise RuntimeError(
                "pyodbc is required for SQL Server connections. Install project dependencies first."
            ) from exc


class ExcelExportBuffer:
    RESULT_BASE_COLUMNS = ["server_ip", "port", "database_name", "scan_status"]
    ERROR_COLUMNS = ["server_ip", "port", "database_name", "error_message"]

    def __init__(self) -> None:
        self._result_columns = list(self.RESULT_BASE_COLUMNS)
        self._results_buffer = tempfile.TemporaryFile(mode="w+", encoding="utf-8")
        self._errors_buffer = tempfile.TemporaryFile(mode="w+", encoding="utf-8")

    def register_query_columns(self, column_names: Sequence[str]) -> list[str]:
        normalized_columns: list[str] = []
        used_columns = set(self.RESULT_BASE_COLUMNS)
        for raw_column_name in column_names:
            base_name = str(raw_column_name).strip() or "column"
            if base_name in self.RESULT_BASE_COLUMNS:
                base_name = f"query_{base_name}"
            candidate_name = base_name
            suffix = 2
            while candidate_name in used_columns:
                candidate_name = f"{base_name}_{suffix}"
                suffix += 1
            used_columns.add(candidate_name)
            normalized_columns.append(candidate_name)
            if candidate_name not in self._result_columns:
                self._result_columns.append(candidate_name)
        return normalized_columns

    def add_result_row(self, row: dict[str, Any]) -> None:
        for column_name in row:
            if column_name not in self._result_columns:
                self._result_columns.append(column_name)
        self._write_json_line(self._results_buffer, row)

    def add_error(self, error: ExportError) -> None:
        self._write_json_line(
            self._errors_buffer,
            {
                "server_ip": error.server_ip,
                "port": error.port,
                "database_name": error.database_name,
                "error_message": error.error_message,
            },
        )

    def save(self, *, output_path: Path) -> None:
        output_path.parent.mkdir(parents=True, exist_ok=True)

        workbook = Workbook(write_only=True)
        results_sheet = workbook.create_sheet(title="results")
        errors_sheet = workbook.create_sheet(title="errors")

        results_sheet.append(self._result_columns)
        self._results_buffer.seek(0)
        for line in self._results_buffer:
            payload = json.loads(line)
            results_sheet.append([payload.get(column_name) for column_name in self._result_columns])

        errors_sheet.append(self.ERROR_COLUMNS)
        self._errors_buffer.seek(0)
        for line in self._errors_buffer:
            payload = json.loads(line)
            errors_sheet.append([payload.get(column_name) for column_name in self.ERROR_COLUMNS])

        workbook.save(output_path)
        self._results_buffer.close()
        self._errors_buffer.close()

    def _write_json_line(self, buffer, payload: dict[str, Any]) -> None:
        json.dump(
            payload,
            buffer,
            ensure_ascii=False,
            default=self._serialize_value,
        )
        buffer.write("\n")

    def _serialize_value(self, value: Any) -> Any:
        if value is None or isinstance(value, (str, int, float, bool)):
            return value
        if isinstance(value, Decimal):
            if value == value.to_integral_value():
                return int(value)
            return float(value)
        if isinstance(value, (datetime, date, time)):
            return value.isoformat(sep=" ")
        if isinstance(value, UUID):
            return str(value)
        if isinstance(value, bytes):
            return value.hex()
        return str(value)


class DatabaseResultExporter:
    def __init__(
        self,
        *,
        input_reader: ServerInputReader,
        connector: SqlServerConnector,
        logger: logging.Logger,
    ) -> None:
        self._input_reader = input_reader
        self._connector = connector
        self._logger = logger

    def run(
        self,
        *,
        input_path: Path,
        query: str,
        output_path: Path,
        allow_write: bool,
        target_database_name: str | None = None,
    ) -> ExportSummary:
        servers, input_errors = self._input_reader.read(input_path=input_path)
        workbook_buffer = ExcelExportBuffer()
        error_count = 0

        for error in input_errors:
            self._logger.error(
                "input_row_invalid server_ip=%s port=%s error=%s",
                error.server_ip,
                error.port,
                error.error_message,
            )
            workbook_buffer.add_error(error)
            error_count += 1

        servers_processed = 0
        databases_scanned = 0
        success_count = 0

        for server in servers:
            servers_processed += 1
            self._logger.info(
                "server_scan_started server_ip=%s port=%s source_row=%s",
                server.host,
                server.port,
                server.source_row,
            )
            if target_database_name is None:
                try:
                    database_names = self._connector.discover_databases(server=server)
                except Exception as exc:
                    error = ExportError(
                        server_ip=server.host,
                        port=server.port,
                        database_name=None,
                        error_message=str(exc),
                    )
                    self._logger.error(
                        "database_discovery_failed server_ip=%s port=%s error=%s",
                        server.host,
                        server.port,
                        exc,
                    )
                    workbook_buffer.add_error(error)
                    error_count += 1
                    continue
            else:
                database_names = [target_database_name]

            self._logger.info(
                "server_databases_discovered server_ip=%s port=%s count=%s",
                server.host,
                server.port,
                len(database_names),
            )
            for current_database_name in database_names:
                databases_scanned += 1
                try:
                    with self._connector.execute_query(
                        server=server,
                        database_name=current_database_name,
                        query=query,
                        allow_write=allow_write,
                    ) as execution_payload:
                        if execution_payload.rows is None:
                            workbook_buffer.add_result_row(
                                {
                                    "server_ip": server.host,
                                    "port": server.port,
                                    "database_name": current_database_name,
                                    "scan_status": SCAN_STATUS_WRITE_OK,
                                    "affected_row_count": execution_payload.affected_row_count,
                                }
                            )
                        else:
                            normalized_columns = workbook_buffer.register_query_columns(
                                execution_payload.columns
                            )
                            row_count = 0
                            for row in execution_payload.rows:
                                row_count += 1
                                payload = {
                                    "server_ip": server.host,
                                    "port": server.port,
                                    "database_name": current_database_name,
                                    "scan_status": SCAN_STATUS_QUERY_OK,
                                }
                                payload.update(
                                    self._build_query_payload(
                                        row=row,
                                        normalized_columns=normalized_columns,
                                    )
                                )
                                workbook_buffer.add_result_row(payload)

                            if row_count == 0:
                                workbook_buffer.add_result_row(
                                    {
                                        "server_ip": server.host,
                                        "port": server.port,
                                        "database_name": current_database_name,
                                        "scan_status": SCAN_STATUS_NO_ROWS,
                                    }
                                )
                        success_count += 1
                except Exception as exc:
                    error = ExportError(
                        server_ip=server.host,
                        port=server.port,
                        database_name=current_database_name,
                        error_message=str(exc),
                    )
                    self._logger.error(
                        "database_query_failed server_ip=%s port=%s database_name=%s error=%s",
                        server.host,
                        server.port,
                        current_database_name,
                        exc,
                    )
                    workbook_buffer.add_error(error)
                    error_count += 1

        workbook_buffer.save(output_path=output_path)
        return ExportSummary(
            servers_processed=servers_processed,
            databases_scanned=databases_scanned,
            success_count=success_count,
            error_count=error_count,
            output_path=output_path,
        )

    def _build_query_payload(
        self,
        *,
        row: Sequence[Any],
        normalized_columns: Sequence[str],
    ) -> dict[str, Any]:
        return {
            normalized_column: row[index]
            for index, normalized_column in enumerate(normalized_columns)
        }


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description=(
            "Scan SQL Server instances, run one query on each non-system database, and export "
            "the results to Excel. Queries are read-only by default."
        ),
    )
    parser.add_argument(
        "--input",
        required=True,
        help="Path to the input server list (.xlsx or .csv).",
    )
    parser.add_argument(
        "--output",
        required=True,
        help="Path to the output Excel workbook.",
    )
    parser.add_argument(
        "--database",
        help=(
            "Optional exact database name to run against on every server. "
            "When omitted, all non-system databases are scanned."
        ),
    )
    query_group = parser.add_mutually_exclusive_group(required=True)
    query_group.add_argument(
        "--query",
        help="Inline SQL query to run on each discovered database.",
    )
    query_group.add_argument(
        "--query-file",
        help="Path to a .sql file containing the query to run on each discovered database.",
    )
    parser.add_argument(
        "--driver",
        default="ODBC Driver 18 for SQL Server",
        help="ODBC driver name for SQL Server connections.",
    )
    parser.add_argument(
        "--connect-timeout",
        type=int,
        default=15,
        help="SQL Server connection timeout in seconds.",
    )
    parser.add_argument(
        "--log-level",
        default="INFO",
        choices=["DEBUG", "INFO", "WARNING", "ERROR"],
        help="Logging level for the command.",
    )
    parser.add_argument(
        "--allow-write",
        action="store_true",
        help=(
            "Allow non-read-only SQL and commit the transaction per database. "
            "Use only for intentional data changes."
        ),
    )
    return parser


def validate_read_only_query(query_text: str) -> str:
    sanitized_query = _sanitize_query_for_validation(query_text)
    tokens = [match.group(0).lower() for match in SQL_TOKEN_PATTERN.finditer(sanitized_query)]
    if not tokens:
        raise ValueError("Query must contain executable SQL.")
    if tokens[0] not in READ_ONLY_ALLOWED_PREFIXES:
        raise ValueError(
            "Only read-only SQL is allowed. Query must start with SELECT, WITH, DECLARE, or SET."
        )

    disallowed_tokens = sorted(
        {
            token
            for token in tokens
            if token in READ_ONLY_DISALLOWED_TOKENS
        }
    )
    if disallowed_tokens:
        raise ValueError(
            "Only read-only SQL is allowed. Disallowed statement(s): "
            + ", ".join(disallowed_tokens)
        )
    return query_text


def _sanitize_query_for_validation(query_text: str) -> str:
    without_block_comments = SQL_BLOCK_COMMENT_PATTERN.sub(" ", query_text)
    without_line_comments = SQL_LINE_COMMENT_PATTERN.sub(" ", without_block_comments)
    return SQL_STRING_LITERAL_PATTERN.sub("''", without_line_comments)


def load_query_text(
    *,
    inline_query: str | None,
    query_file: str | None,
    allow_write: bool,
) -> str:
    if inline_query:
        query_text = inline_query.strip()
        if not query_text:
            raise ValueError("Query text must not be empty.")
        if allow_write:
            return query_text
        return validate_read_only_query(query_text)
    if query_file:
        query_text = Path(query_file).read_text(encoding="utf-8").strip()
        if not query_text:
            raise ValueError("Query file is empty.")
        if allow_write:
            return query_text
        return validate_read_only_query(query_text)
    raise ValueError("Either --query or --query-file must be provided.")


def configure_logging(*, log_level: str) -> logging.Logger:
    logging.basicConfig(
        level=getattr(logging, log_level.upper()),
        format="%(asctime)s %(levelname)s %(message)s",
    )
    return logging.getLogger("db_results_export")


def format_summary(summary: ExportSummary) -> str:
    return (
        "servers_processed="
        f"{summary.servers_processed}\n"
        f"databases_scanned={summary.databases_scanned}\n"
        f"success_count={summary.success_count}\n"
        f"error_count={summary.error_count}\n"
        f"output_file={summary.output_path}"
    )


def main(argv: Sequence[str] | None = None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)
    logger = configure_logging(log_level=args.log_level)

    query_text = load_query_text(
        inline_query=args.query,
        query_file=args.query_file,
        allow_write=args.allow_write,
    )
    exporter = DatabaseResultExporter(
        input_reader=ServerInputReader(),
        connector=SqlServerConnector(
            driver=args.driver,
            connect_timeout=args.connect_timeout,
        ),
        logger=logger,
    )
    summary = exporter.run(
        input_path=Path(args.input).resolve(),
        query=query_text,
        output_path=Path(args.output).resolve(),
        allow_write=args.allow_write,
        target_database_name=args.database,
    )
    summary_text = format_summary(summary)
    logger.info("export_completed\n%s", summary_text)
    print(summary_text)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
