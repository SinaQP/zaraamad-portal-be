from __future__ import annotations

import json
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from datetime import datetime, timezone
from pathlib import Path

from fastapi import Depends, HTTPException, status
from openpyxl import load_workbook
from sqlalchemy import Select, delete, func, or_, select
from sqlalchemy.engine import make_url
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.common.config import Settings, get_settings
from app.common.database import get_db_session
from app.common.dtos import CurrentUser
from app.common.enums import SortOrder, UserRole
from app.common.messages import (
    CUSTOMER_ACCESS_DENIED,
    CUSTOMER_BRIDGE_AUTH_FAILED,
    CUSTOMER_BRIDGE_INVALID_RESPONSE,
    CUSTOMER_BRIDGE_NOT_CONFIGURED,
    CUSTOMER_BRIDGE_REQUEST_FAILED,
    CUSTOMER_BRIDGE_SUBSCRIPTION_NOT_CACHED,
    CUSTOMER_BRIDGE_UNAVAILABLE,
    CUSTOMER_DATABASE_CONNECTION_INCOMPLETE,
    CUSTOMER_DATABASE_CONNECTION_NOT_FOUND,
    CUSTOMER_DATABASE_CONNECTION_SECRET_NOT_CONFIGURED,
    CUSTOMER_DATABASE_CONNECTION_STRING_INVALID,
    CUSTOMER_DATABASE_CONNECTION_TEST_FAILED,
    CUSTOMER_ID_INVALID_OR_INACTIVE,
    CUSTOMER_INCOME_BUCKET_TOTAL_MISMATCH,
    CUSTOMER_NOT_FOUND,
    DATA_INTEGRITY_ERROR,
    DUPLICATE_CUSTOMER_INCOME_BUCKET_CODE,
    DUPLICATE_CUSTOMER_INCOME_CUSTOMER,
    DUPLICATE_CUSTOMER_INCOME_REPORT_MONTH,
)
from app.common.security.secret_cipher import (
    SecretCipherConfigurationError,
    SecretCipherDecryptionError,
    SecretCipherService,
    get_secret_cipher_service,
)
from app.common.security.bridge_token_service import (
    BridgeAccessTokenService,
    get_bridge_access_token_service,
)
from app.common.pagination import PaginationMeta, PaginationParams
from app.common.services.bridge_client import (
    BridgeCapabilitiesResult,
    BridgeClient,
    BridgeConnectionError,
    BridgeHealthResult,
    BridgeSubscriptionConfigResult,
    BridgeSubscriptionMessageResult,
    BridgeSubscriptionResult,
    BridgeInvalidResponseError,
    BridgeRequest,
    BridgeUnauthorizedError,
    BridgeUnexpectedStatusError,
    get_bridge_client,
)
from app.common.services.sqlserver_reference_sync import (
    SqlServerConnectionError,
    SqlServerConnectionSettings,
    SqlServerEngineFactory,
)
from app.modules.customers.dtos import (
    CustomerBridgeConfigUpdate,
    CustomerCreate,
    CustomerDatabaseConnectionCreate,
    CustomerDatabaseConnectionUpdate,
    CustomerIncomeBulkUpsertCreate,
    CustomerIncomeBulkUpsertItemCreate,
    CustomerUpdate,
)
from app.modules.customers.schemas import (
    Customer,
    CustomerBridgeConfig,
    CustomerDatabaseConnection,
    CustomerIncomeBucket,
    CustomerIncomeMonthlyReport,
    CustomerIncomeSummary,
)


class CustomerQueryBuilder:
    SORT_COLUMNS = {
        "id": Customer.id,
        "name": Customer.name,
        "manager_name": Customer.manager_name,
        "grade": Customer.grade,
        "is_active": Customer.is_active,
        "created_at": Customer.created_at,
        "updated_at": Customer.updated_at,
    }

    def build_list_query(
        self,
        is_active: bool | None,
        search: str | None,
        sort_by: str,
        sort_order: SortOrder,
    ) -> Select[tuple[Customer]]:
        query = select(Customer)
        if is_active is None:
            query = query.where(Customer.is_active.is_(True))
        else:
            query = query.where(Customer.is_active.is_(is_active))
        if search:
            search_pattern = f"%{search}%"
            query = query.where(
                or_(
                    Customer.name.ilike(search_pattern),
                    Customer.manager_name.ilike(search_pattern),
                )
            )
        sort_column = self.SORT_COLUMNS[sort_by]
        if sort_order == SortOrder.DESC:
            query = query.order_by(sort_column.desc(), Customer.id.desc())
        else:
            query = query.order_by(sort_column.asc(), Customer.id.asc())
        return query

    def build_without_income_query(
        self,
        *,
        search: str | None,
        sort_by: str,
        sort_order: SortOrder,
    ) -> Select[tuple[Customer]]:
        query = (
            select(Customer)
            .outerjoin(CustomerIncomeSummary, CustomerIncomeSummary.customer_id == Customer.id)
            .where(Customer.is_active.is_(True))
            .where(CustomerIncomeSummary.customer_id.is_(None))
        )
        if search:
            search_pattern = f"%{search}%"
            query = query.where(
                or_(
                    Customer.name.ilike(search_pattern),
                    Customer.manager_name.ilike(search_pattern),
                )
            )
        sort_column = self.SORT_COLUMNS[sort_by]
        if sort_order == SortOrder.DESC:
            query = query.order_by(sort_column.desc(), Customer.id.desc())
        else:
            query = query.order_by(sort_column.asc(), Customer.id.asc())
        return query


class CustomerService:
    def __init__(self, db_session: Session) -> None:
        self._db_session = db_session
        self._query_builder = CustomerQueryBuilder()

    def create(self, dto: CustomerCreate) -> Customer:
        customer = Customer(
            name=dto.name,
            manager_name=dto.manager_name,
            grade=dto.grade,
            is_active=True,
        )
        self._db_session.add(customer)
        try:
            self._db_session.commit()
        except IntegrityError as exc:
            self._db_session.rollback()
            raise HTTPException(
                status_code=status.HTTP_409_CONFLICT,
                detail=DATA_INTEGRITY_ERROR,
            ) from exc
        self._db_session.refresh(customer)
        return customer

    def list_customers(
        self,
        is_active: bool | None,
        search: str | None,
        sort_by: str,
        sort_order: SortOrder,
        pagination: PaginationParams,
    ) -> tuple[list[Customer], PaginationMeta]:
        query = self._query_builder.build_list_query(
            is_active=is_active,
            search=search,
            sort_by=sort_by,
            sort_order=sort_order,
        )
        total_count = int(
            self._db_session.scalar(
                select(func.count()).select_from(query.order_by(None).subquery())
            ) or 0
        )
        paginated_query = query.offset(pagination.offset).limit(pagination.page_size)
        items = list(self._db_session.scalars(paginated_query).all())
        meta = PaginationMeta(
            total_count=total_count,
            page=pagination.page,
            page_size=pagination.page_size,
        )
        return items, meta

    def list_all(self) -> list[Customer]:
        query = (
            select(Customer)
            .where(Customer.is_active.is_(True))
            .order_by(Customer.id.asc())
        )
        return list(self._db_session.scalars(query).all())

    def list_without_income(
        self,
        *,
        search: str | None,
        sort_by: str,
        sort_order: SortOrder,
        pagination: PaginationParams,
    ) -> tuple[list[Customer], PaginationMeta]:
        query = self._query_builder.build_without_income_query(
            search=search,
            sort_by=sort_by,
            sort_order=sort_order,
        )
        total_count = int(
            self._db_session.scalar(
                select(func.count()).select_from(query.order_by(None).subquery())
            ) or 0
        )
        paginated_query = query.offset(pagination.offset).limit(pagination.page_size)
        items = list(self._db_session.scalars(paginated_query).all())
        meta = PaginationMeta(
            total_count=total_count,
            page=pagination.page,
            page_size=pagination.page_size,
        )
        return items, meta

    def get_or_404(self, customer_id: int) -> Customer:
        customer = self._db_session.get(Customer, customer_id)
        if customer is None:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=CUSTOMER_NOT_FOUND,
            )
        return customer

    def get_active_or_404(self, customer_id: int) -> Customer:
        customer = self.get_or_404(customer_id=customer_id)
        if not customer.is_active:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=CUSTOMER_NOT_FOUND,
            )
        return customer

    def update(self, customer_id: int, dto: CustomerUpdate) -> Customer:
        customer = self.get_or_404(customer_id=customer_id)
        update_data = dto.model_dump(exclude_unset=True, exclude_none=False)
        for field_name, field_value in update_data.items():
            setattr(customer, field_name, field_value)
        try:
            self._db_session.commit()
        except IntegrityError as exc:
            self._db_session.rollback()
            raise HTTPException(
                status_code=status.HTTP_409_CONFLICT,
                detail=DATA_INTEGRITY_ERROR,
            ) from exc
        self._db_session.refresh(customer)
        return customer

    def deactivate(self, customer_id: int) -> Customer:
        customer = self.get_or_404(customer_id=customer_id)
        customer.is_active = False
        self._db_session.commit()
        self._db_session.refresh(customer)
        return customer


class CustomerBridgeConfigService:
    def __init__(
        self,
        db_session: Session,
        customer_service: CustomerService,
    ) -> None:
        self._db_session = db_session
        self._customer_service = customer_service

    def get(
        self,
        customer_id: int,
        instance_id: str | None = None,
    ) -> tuple[Customer, CustomerBridgeConfig | None]:
        customer = self._customer_service.get_or_404(customer_id=customer_id)
        bridge_config = self._find_bridge_config(
            customer_id=customer_id,
            instance_id=instance_id,
        )
        return customer, bridge_config

    def get_active(
        self,
        customer_id: int,
        instance_id: str | None = None,
    ) -> tuple[Customer, CustomerBridgeConfig | None]:
        customer = self._customer_service.get_active_or_404(customer_id=customer_id)
        bridge_config = self._find_bridge_config(
            customer_id=customer_id,
            instance_id=instance_id,
        )
        return customer, bridge_config

    def upsert(
        self,
        customer_id: int,
        dto: CustomerBridgeConfigUpdate,
    ) -> tuple[Customer, CustomerBridgeConfig]:
        customer = self._customer_service.get_or_404(customer_id=customer_id)
        bridge_config = self._find_bridge_config(
            customer_id=customer_id,
            instance_id=None,
        )
        requested_instance_id = dto.instance_id or "default"
        if bridge_config is None:
            bridge_config = CustomerBridgeConfig(
                customer_id=customer_id,
                instance_id=requested_instance_id,
            )
            self._db_session.add(bridge_config)
        elif dto.instance_id is not None:
            bridge_config.instance_id = dto.instance_id
        update_data = dto.model_dump(exclude_unset=True, exclude_none=False)
        for field_name, field_value in update_data.items():
            setattr(bridge_config, field_name, field_value)
        if not bridge_config.instance_id:
            bridge_config.instance_id = requested_instance_id
        if not bridge_config.audience:
            bridge_config.audience = f"zaraamad:{bridge_config.instance_id}"
        if not bridge_config.tenant_id:
            bridge_config.tenant_id = str(customer_id)
        if not bridge_config.status:
            bridge_config.status = "inactive"
        if update_data:
            self.clear_health_status(bridge_config=bridge_config)
            self.clear_subscription_cache(bridge_config=bridge_config)
        try:
            self._db_session.commit()
        except IntegrityError as exc:
            self._db_session.rollback()
            raise HTTPException(
                status_code=status.HTTP_409_CONFLICT,
                detail=DATA_INTEGRITY_ERROR,
            ) from exc
        self._db_session.refresh(bridge_config)
        return customer, bridge_config

    def clear_health_status(self, bridge_config: CustomerBridgeConfig) -> None:
        bridge_config.last_online_status = None
        bridge_config.last_health_checked_at = None
        bridge_config.last_health_error = None

    def clear_subscription_cache(self, bridge_config: CustomerBridgeConfig) -> None:
        bridge_config.cached_subscription_start_date = None
        bridge_config.cached_subscription_end_date = None
        bridge_config.cached_subscription_grace_period_end_date = None
        bridge_config.cached_subscription_is_active = None
        bridge_config.cached_subscription_status_message = None
        bridge_config.last_subscription_synced_at = None
        bridge_config.last_subscription_error = None

    def persist_health_status(
        self,
        bridge_config: CustomerBridgeConfig,
        *,
        is_online: bool,
        checked_at: datetime,
        error: str | None,
    ) -> CustomerBridgeConfig:
        bridge_config.last_online_status = is_online
        bridge_config.last_health_checked_at = checked_at
        bridge_config.last_health_error = error
        self._db_session.commit()
        self._db_session.refresh(bridge_config)
        return bridge_config

    def persist_subscription_snapshot(
        self,
        bridge_config: CustomerBridgeConfig,
        *,
        subscription: BridgeSubscriptionResult,
        checked_at: datetime,
    ) -> CustomerBridgeConfig:
        bridge_config.cached_subscription_start_date = subscription.start_date
        bridge_config.cached_subscription_end_date = subscription.end_date
        bridge_config.cached_subscription_grace_period_end_date = subscription.grace_period_end_date
        bridge_config.cached_subscription_is_active = subscription.is_active
        bridge_config.cached_subscription_status_message = subscription.status_message
        bridge_config.last_subscription_synced_at = checked_at
        bridge_config.last_subscription_error = None
        self._db_session.commit()
        self._db_session.refresh(bridge_config)
        return bridge_config

    def persist_subscription_absence(
        self,
        bridge_config: CustomerBridgeConfig,
        *,
        checked_at: datetime,
        error: str,
    ) -> CustomerBridgeConfig:
        self.clear_subscription_cache(bridge_config=bridge_config)
        bridge_config.last_subscription_synced_at = checked_at
        bridge_config.last_subscription_error = error
        self._db_session.commit()
        self._db_session.refresh(bridge_config)
        return bridge_config

    def persist_subscription_refresh_error(
        self,
        bridge_config: CustomerBridgeConfig,
        *,
        error: str,
    ) -> CustomerBridgeConfig:
        bridge_config.last_subscription_error = error
        self._db_session.commit()
        self._db_session.refresh(bridge_config)
        return bridge_config

    def list_refreshable_customer_ids(self) -> list[int]:
        query = (
            select(CustomerBridgeConfig.customer_id).distinct()
            .where(CustomerBridgeConfig.status == "active")
            .where(CustomerBridgeConfig.base_url_internal.is_not(None))
            .where(CustomerBridgeConfig.base_url_internal != "")
            .where(CustomerBridgeConfig.instance_id.is_not(None))
            .where(CustomerBridgeConfig.instance_id != "")
            .where(CustomerBridgeConfig.audience.is_not(None))
            .where(CustomerBridgeConfig.audience != "")
            .where(CustomerBridgeConfig.tenant_id.is_not(None))
            .where(CustomerBridgeConfig.tenant_id != "")
            .order_by(CustomerBridgeConfig.customer_id.asc())
        )
        return list(self._db_session.scalars(query).all())

    def _find_bridge_config(
        self,
        *,
        customer_id: int,
        instance_id: str | None,
    ) -> CustomerBridgeConfig | None:
        query = (
            select(CustomerBridgeConfig)
            .where(CustomerBridgeConfig.customer_id == customer_id)
            .order_by(CustomerBridgeConfig.updated_at.desc())
        )
        if instance_id is not None:
            query = query.where(CustomerBridgeConfig.instance_id == instance_id)
        results = list(self._db_session.scalars(query).all())
        if len(results) > 1 and instance_id is None:
            raise HTTPException(
                status_code=status.HTTP_409_CONFLICT,
                detail={
                    "message": CUSTOMER_BRIDGE_NOT_CONFIGURED,
                    "developer_message": (
                        f"Customer {customer_id} has multiple bridge instances. "
                        "Provide instance_id to resolve the target Bridge row."
                    ),
                },
            )
        return results[0] if results else None

class CustomerDatabaseConnectionValidationError(Exception):
    pass


class CustomerDatabaseConnectionRuntimeError(Exception):
    pass


class CustomerDatabaseConnectionVerifier:
    DEFAULT_DRIVER_NAME = "ODBC Driver 18 for SQL Server"
    DEFAULT_PORT = 1433
    DEFAULT_CONNECT_TIMEOUT_SECONDS = 15

    def __init__(self) -> None:
        self._engine_factory = SqlServerEngineFactory(
            driver=self.DEFAULT_DRIVER_NAME,
            connect_timeout=self.DEFAULT_CONNECT_TIMEOUT_SECONDS,
        )

    def build_settings(self, *, connection_string: str) -> SqlServerConnectionSettings:
        normalized_connection_string = connection_string.strip()
        if not normalized_connection_string:
            raise CustomerDatabaseConnectionValidationError(
                "Customer database connection string is empty."
            )
        try:
            connection_url = make_url(normalized_connection_string)
        except Exception as exc:
            raise CustomerDatabaseConnectionValidationError(
                "Customer database connection string could not be parsed."
            ) from exc
        if connection_url.drivername != "mssql+pyodbc":
            raise CustomerDatabaseConnectionValidationError(
                "Customer database connection must use the mssql+pyodbc driver."
            )
        if not connection_url.host:
            raise CustomerDatabaseConnectionValidationError(
                "Customer database connection host is required."
            )
        if not connection_url.username:
            raise CustomerDatabaseConnectionValidationError(
                "Customer database connection username is required."
            )
        if connection_url.password is None or not str(connection_url.password).strip():
            raise CustomerDatabaseConnectionValidationError(
                "Customer database connection password is required."
            )
        if not connection_url.database:
            raise CustomerDatabaseConnectionValidationError(
                "Customer database connection database name is required."
            )
        query_lookup = {
            str(key).strip().lower(): str(value).strip()
            for key, value in connection_url.query.items()
        }
        return SqlServerConnectionSettings(
            host=connection_url.host,
            port=connection_url.port or self.DEFAULT_PORT,
            username=connection_url.username,
            password=str(connection_url.password),
            database_name=connection_url.database,
            driver_name=query_lookup.get("driver") or self.DEFAULT_DRIVER_NAME,
            encrypt_connection=self._parse_connection_flag(
                raw_value=query_lookup.get("encrypt"),
                default=True,
            ),
            trust_server_certificate=self._parse_connection_flag(
                raw_value=query_lookup.get("trustservercertificate"),
                default=False,
            ),
        )

    def test_connection(self, *, connection_string: str) -> None:
        settings = self.build_settings(connection_string=connection_string)
        try:
            self._engine_factory._check_connection(settings=settings)
        except SqlServerConnectionError as exc:
            raise CustomerDatabaseConnectionRuntimeError(
                "Customer database connection test failed."
            ) from exc

    def _parse_connection_flag(self, *, raw_value: str | None, default: bool) -> bool:
        if raw_value is None:
            return default
        normalized_value = raw_value.strip().lower()
        if normalized_value in {"1", "true", "yes"}:
            return True
        if normalized_value in {"0", "false", "no"}:
            return False
        return default


class CustomerDatabaseConnectionService:
    SECRET_MASK_TOKENS = ("password", "pwd", "secret", "token")

    def __init__(
        self,
        db_session: Session,
        customer_service: CustomerService,
        secret_cipher_service: SecretCipherService,
        verifier: CustomerDatabaseConnectionVerifier,
    ) -> None:
        self._db_session = db_session
        self._customer_service = customer_service
        self._secret_cipher_service = secret_cipher_service
        self._verifier = verifier

    def get(self, customer_id: int) -> tuple[Customer, CustomerDatabaseConnection | None]:
        customer = self._customer_service.get_or_404(customer_id=customer_id)
        database_connection = self._db_session.get(CustomerDatabaseConnection, customer_id)
        return customer, database_connection

    def get_active(self, customer_id: int) -> tuple[Customer, CustomerDatabaseConnection]:
        customer = self._customer_service.get_active_or_404(customer_id=customer_id)
        database_connection = self._db_session.get(CustomerDatabaseConnection, customer_id)
        if database_connection is None or not database_connection.is_active:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=CUSTOMER_DATABASE_CONNECTION_NOT_FOUND,
            )
        return customer, database_connection

    def upsert(
        self,
        customer_id: int,
        dto: CustomerDatabaseConnectionCreate | CustomerDatabaseConnectionUpdate,
    ) -> tuple[Customer, CustomerDatabaseConnection]:
        customer, database_connection = self.get(customer_id=customer_id)
        if database_connection is None:
            database_connection = CustomerDatabaseConnection(
                customer_id=customer_id,
                db_kind="sqlserver",
            )
            self._db_session.add(database_connection)
        update_data = dto.model_dump(exclude_unset=True, exclude_none=False)
        connection_string = update_data.pop("connection_string", None)
        for field_name, field_value in update_data.items():
            setattr(database_connection, field_name, field_value)
        if connection_string is not None:
            self._apply_encrypted_connection_string(
                database_connection=database_connection,
                connection_string=connection_string,
            )
            if "credential_rotated_at" not in update_data:
                database_connection.credential_rotated_at = datetime.now(timezone.utc)
        database_connection.last_connection_test_success = None
        database_connection.last_connection_tested_at = None
        database_connection.last_connection_error = None
        try:
            self._db_session.commit()
        except IntegrityError as exc:
            self._db_session.rollback()
            raise HTTPException(
                status_code=status.HTTP_409_CONFLICT,
                detail=DATA_INTEGRITY_ERROR,
            ) from exc
        self._db_session.refresh(database_connection)
        return customer, database_connection

    def mark_rotation(
        self,
        *,
        customer_id: int,
        rotated_at: datetime,
        next_rotation_due_at: datetime | None,
        secret_version: str | None,
    ) -> tuple[Customer, CustomerDatabaseConnection]:
        customer, database_connection = self.get_active(customer_id=customer_id)
        database_connection.credential_rotated_at = rotated_at
        database_connection.rotation_due_at = next_rotation_due_at
        database_connection.secret_version = secret_version
        database_connection.last_connection_test_success = None
        database_connection.last_connection_tested_at = None
        database_connection.last_connection_error = None
        self._db_session.commit()
        self._db_session.refresh(database_connection)
        return customer, database_connection

    def persist_connection_test_result(
        self,
        *,
        customer_id: int,
        checked_at: datetime,
        is_success: bool,
        error_message: str | None = None,
    ) -> tuple[Customer, CustomerDatabaseConnection]:
        customer, database_connection = self.get_active(customer_id=customer_id)
        database_connection.last_connection_tested_at = checked_at
        database_connection.last_connection_test_success = is_success
        database_connection.last_connection_error = (
            None if is_success else self._sanitize_error_message(error_message)
        )
        self._db_session.commit()
        self._db_session.refresh(database_connection)
        return customer, database_connection

    def build_sqlserver_connection_settings(
        self,
        *,
        customer_id: int,
    ) -> tuple[Customer, CustomerDatabaseConnection, SqlServerConnectionSettings]:
        customer, database_connection = self.get_active(customer_id=customer_id)
        decrypted_connection_string = self._decrypt_connection_string(
            database_connection=database_connection,
        )
        settings = self._verifier.build_settings(
            connection_string=decrypted_connection_string,
        )
        return customer, database_connection, settings

    def test_connection(
        self,
        *,
        customer_id: int,
    ) -> tuple[Customer, CustomerDatabaseConnection]:
        customer, database_connection = self.get_active(customer_id=customer_id)
        decrypted_connection_string = self._decrypt_connection_string(
            database_connection=database_connection,
        )
        checked_at = datetime.now(timezone.utc)
        try:
            self._verifier.test_connection(
                connection_string=decrypted_connection_string,
            )
        except CustomerDatabaseConnectionValidationError as exc:
            self.persist_connection_test_result(
                customer_id=customer_id,
                checked_at=checked_at,
                is_success=False,
                error_message=str(exc),
            )
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail={
                    "message": CUSTOMER_DATABASE_CONNECTION_STRING_INVALID,
                    "developer_message": str(exc),
                },
            ) from exc
        except CustomerDatabaseConnectionRuntimeError as exc:
            self.persist_connection_test_result(
                customer_id=customer_id,
                checked_at=checked_at,
                is_success=False,
                error_message=str(exc),
            )
            raise HTTPException(
                status_code=status.HTTP_502_BAD_GATEWAY,
                detail={
                    "message": CUSTOMER_DATABASE_CONNECTION_TEST_FAILED,
                    "developer_message": self._sanitize_error_message(str(exc))
                    or "Customer database connection test failed.",
                },
            ) from exc
        return self.persist_connection_test_result(
            customer_id=customer_id,
            checked_at=checked_at,
            is_success=True,
            error_message=None,
        )

    def _apply_encrypted_connection_string(
        self,
        *,
        database_connection: CustomerDatabaseConnection,
        connection_string: str,
    ) -> None:
        try:
            settings = self._verifier.build_settings(
                connection_string=connection_string,
            )
            encrypted_connection_string = self._secret_cipher_service.encrypt(
                connection_string.strip(),
            )
            connection_string_hash = self._secret_cipher_service.fingerprint(
                connection_string.strip(),
            )
        except CustomerDatabaseConnectionValidationError as exc:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail={
                    "message": CUSTOMER_DATABASE_CONNECTION_STRING_INVALID,
                    "developer_message": str(exc),
                },
            ) from exc
        except SecretCipherConfigurationError as exc:
            raise HTTPException(
                status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
                detail={
                    "message": CUSTOMER_DATABASE_CONNECTION_SECRET_NOT_CONFIGURED,
                    "developer_message": str(exc),
                },
            ) from exc
        database_connection.db_kind = "sqlserver"
        database_connection.encrypted_connection_string = encrypted_connection_string
        database_connection.connection_string_hash = connection_string_hash
        database_connection.host = None
        database_connection.port = None
        database_connection.database_name = None
        database_connection.username = None
        database_connection.secret_ref = None
        database_connection.driver_name = settings.driver_name
        database_connection.encrypt_connection = settings.encrypt_connection
        database_connection.trust_server_certificate = settings.trust_server_certificate

    def _decrypt_connection_string(
        self,
        *,
        database_connection: CustomerDatabaseConnection,
    ) -> str:
        if not database_connection.encrypted_connection_string:
            if database_connection.secret_ref:
                raise HTTPException(
                    status_code=status.HTTP_409_CONFLICT,
                    detail={
                        "message": CUSTOMER_DATABASE_CONNECTION_INCOMPLETE,
                        "developer_message": (
                            "Legacy customer database connection fields exist without "
                            "an encrypted connection string."
                        ),
                    },
                )
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=CUSTOMER_DATABASE_CONNECTION_NOT_FOUND,
            )
        try:
            return self._secret_cipher_service.decrypt(
                database_connection.encrypted_connection_string,
            )
        except SecretCipherConfigurationError as exc:
            raise HTTPException(
                status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
                detail={
                    "message": CUSTOMER_DATABASE_CONNECTION_SECRET_NOT_CONFIGURED,
                    "developer_message": str(exc),
                },
            ) from exc
        except SecretCipherDecryptionError as exc:
            raise HTTPException(
                status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
                detail={
                    "message": CUSTOMER_DATABASE_CONNECTION_INCOMPLETE,
                    "developer_message": str(exc),
                },
            ) from exc

    def _sanitize_error_message(self, error_message: str | None) -> str | None:
        if error_message is None:
            return None
        normalized_error_message = error_message.strip()
        if not normalized_error_message:
            return None
        sanitized_message = normalized_error_message
        for token in self.SECRET_MASK_TOKENS:
            sanitized_message = sanitized_message.replace(token, "***")
            sanitized_message = sanitized_message.replace(token.upper(), "***")
            sanitized_message = sanitized_message.replace(token.capitalize(), "***")
        return sanitized_message[:1000]


DEFAULT_IMPORTED_CUSTOMER_GRADE = 1


@dataclass(frozen=True)
class ImportedCustomerIncomeSummary:
    customer_name: str
    registered_income_amount: int | None
    issued_bill_count: int | None
    paid_bill_count: int | None
    collection_rate_percent: float | None


@dataclass(frozen=True)
class ImportedCustomerIncomeBucket:
    customer_name: str
    bucket_code: str
    bucket_name: str | None
    registered_income_amount: int | None


@dataclass(frozen=True)
class CustomerIncomeValidationFailure:
    customer_name: str
    message: str
    summary_registered_income_amount: int | None
    bucket_total_registered_income_amount: int | None


@dataclass(frozen=True)
class CustomerIncomeWorkbookData:
    summaries: list[ImportedCustomerIncomeSummary]
    buckets: list[ImportedCustomerIncomeBucket]
    validation_failures: list[CustomerIncomeValidationFailure]


@dataclass(frozen=True)
class CustomerIncomeImportReport:
    processed_customer_count: int
    upserted_summary_count: int
    upserted_bucket_count: int
    validation_failures: list[CustomerIncomeValidationFailure]


class CustomerIncomeQueryBuilder:
    SORT_COLUMNS = {
        "customer_name": Customer.name,
        "registered_income_amount": CustomerIncomeSummary.registered_income_amount,
        "issued_bill_count": CustomerIncomeSummary.issued_bill_count,
        "paid_bill_count": CustomerIncomeSummary.paid_bill_count,
        "collection_rate_percent": CustomerIncomeSummary.collection_rate_percent,
        "created_at": CustomerIncomeSummary.created_at,
        "updated_at": CustomerIncomeSummary.updated_at,
    }

    def build_list_query(
        self,
        *,
        search: str | None,
        sort_by: str,
        sort_order: SortOrder,
    ) -> Select[tuple[Customer, CustomerIncomeSummary]]:
        query = (
            select(Customer, CustomerIncomeSummary)
            .join(CustomerIncomeSummary, CustomerIncomeSummary.customer_id == Customer.id)
            .where(Customer.is_active.is_(True))
        )
        if search:
            search_pattern = f"%{search}%"
            query = query.where(Customer.name.ilike(search_pattern))
        sort_column = self.SORT_COLUMNS[sort_by]
        if sort_order == SortOrder.DESC:
            query = query.order_by(sort_column.desc(), Customer.id.desc())
        else:
            query = query.order_by(sort_column.asc(), Customer.id.asc())
        return query

class CustomerScopedAccessPolicy:
    def validate_customer_access(
        self,
        *,
        current_user: CurrentUser,
        customer_id: int,
    ) -> None:
        role_label = current_user.role.value if current_user.role is not None else "unknown"
        if current_user.role == UserRole.ADMIN:
            return
        has_customer_access = (
            current_user.role == UserRole.CUSTOMER
            and current_user.customer_id == customer_id
        )
        if has_customer_access:
            return
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail={
                "message": CUSTOMER_ACCESS_DENIED,
                "developer_message": (
                    f"User {current_user.user_id} with role {role_label} "
                    f"cannot access customer {customer_id}."
                ),
            },
        )


class CustomerIncomeMutationPolicy:
    def __init__(self, db_session: Session) -> None:
        self._db_session = db_session

    def validate_payload(self, *, items: list[CustomerIncomeBulkUpsertItemCreate]) -> None:
        self._validate_unique_customer_ids(items=items)
        for item in items:
            self._validate_unique_bucket_codes(item=item)
            self._validate_unique_report_months(item=item)
            self._validate_bucket_total(item=item)

    def get_active_customer_map(
        self,
        *,
        customer_ids: list[int],
    ) -> dict[int, Customer]:
        customers = list(
            self._db_session.scalars(
                select(Customer)
                .where(Customer.id.in_(customer_ids))
                .where(Customer.is_active.is_(True))
            ).all()
        )
        customer_map = {customer.id: customer for customer in customers}
        if len(customer_map) != len(set(customer_ids)):
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail=CUSTOMER_ID_INVALID_OR_INACTIVE,
            )
        return customer_map

    def _validate_unique_customer_ids(
        self,
        *,
        items: list[CustomerIncomeBulkUpsertItemCreate],
    ) -> None:
        customer_ids = [item.customer_id for item in items]
        if len(customer_ids) != len(set(customer_ids)):
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail=DUPLICATE_CUSTOMER_INCOME_CUSTOMER,
            )

    def _validate_unique_bucket_codes(
        self,
        *,
        item: CustomerIncomeBulkUpsertItemCreate,
    ) -> None:
        bucket_codes = [bucket.bucket_code for bucket in item.buckets]
        if len(bucket_codes) != len(set(bucket_codes)):
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail=DUPLICATE_CUSTOMER_INCOME_BUCKET_CODE,
            )

    def _validate_bucket_total(
        self,
        *,
        item: CustomerIncomeBulkUpsertItemCreate,
    ) -> None:
        summary_amount = item.summary.registered_income_amount
        bucket_amounts = [
            bucket.registered_income_amount
            for bucket in item.buckets
            if bucket.registered_income_amount is not None
        ]
        has_bucket_values = len(bucket_amounts) > 0
        bucket_total = sum(bucket_amounts)
        if summary_amount is None and not has_bucket_values:
            return
        if summary_amount != bucket_total:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail=CUSTOMER_INCOME_BUCKET_TOTAL_MISMATCH,
            )

    def _validate_unique_report_months(
        self,
        *,
        item: CustomerIncomeBulkUpsertItemCreate,
    ) -> None:
        if item.monthly_reports is None:
            return
        report_months = [report.month for report in item.monthly_reports]
        if len(report_months) != len(set(report_months)):
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail=DUPLICATE_CUSTOMER_INCOME_REPORT_MONTH,
            )


class CustomerIncomeService:
    def __init__(
        self,
        db_session: Session,
        customer_service: CustomerService,
    ) -> None:
        self._db_session = db_session
        self._customer_service = customer_service
        self._query_builder = CustomerIncomeQueryBuilder()
        self._access_policy = CustomerScopedAccessPolicy()
        self._mutation_policy = CustomerIncomeMutationPolicy(db_session=db_session)

    def list_summaries(
        self,
        *,
        search: str | None,
        sort_by: str,
        sort_order: SortOrder,
        pagination: PaginationParams,
    ) -> tuple[list[tuple[Customer, CustomerIncomeSummary]], PaginationMeta]:
        query = self._query_builder.build_list_query(
            search=search,
            sort_by=sort_by,
            sort_order=sort_order,
        )
        total_count = int(
            self._db_session.scalar(
                select(func.count()).select_from(query.order_by(None).subquery())
            ) or 0
        )
        paginated_query = query.offset(pagination.offset).limit(pagination.page_size)
        items = list(self._db_session.execute(paginated_query).all())
        meta = PaginationMeta(
            total_count=total_count,
            page=pagination.page,
            page_size=pagination.page_size,
        )
        return items, meta

    def get_detail(
        self,
        *,
        customer_id: int,
        current_user: CurrentUser,
    ) -> tuple[
        Customer,
        CustomerIncomeSummary | None,
        list[CustomerIncomeBucket],
        list[CustomerIncomeMonthlyReport],
    ]:
        customer = self._customer_service.get_active_or_404(customer_id=customer_id)
        self._access_policy.validate_customer_access(
            current_user=current_user,
            customer_id=customer.id,
        )
        income_summary = self._db_session.get(CustomerIncomeSummary, customer.id)
        bucket_query = (
            select(CustomerIncomeBucket)
            .where(CustomerIncomeBucket.customer_id == customer.id)
            .order_by(
                CustomerIncomeBucket.bucket_code.asc(),
                CustomerIncomeBucket.id.asc(),
            )
        )
        buckets = list(self._db_session.scalars(bucket_query).all())
        monthly_reports = list(
            self._db_session.scalars(
                select(CustomerIncomeMonthlyReport)
                .where(CustomerIncomeMonthlyReport.customer_id == customer.id)
                .order_by(
                    CustomerIncomeMonthlyReport.month.asc(),
                    CustomerIncomeMonthlyReport.id.asc(),
                )
            ).all()
        )
        return customer, income_summary, buckets, monthly_reports

    def bulk_upsert(
        self,
        *,
        dto: CustomerIncomeBulkUpsertCreate,
    ) -> list[
        tuple[
            Customer,
            CustomerIncomeSummary,
            list[CustomerIncomeBucket],
            list[CustomerIncomeMonthlyReport],
        ]
    ]:
        self._mutation_policy.validate_payload(items=dto.items)
        customer_ids = [item.customer_id for item in dto.items]
        customer_map = self._mutation_policy.get_active_customer_map(customer_ids=customer_ids)
        summary_map = self._get_income_summary_map(customer_ids=customer_ids)
        monthly_report_customer_ids = [
            item.customer_id
            for item in dto.items
            if item.monthly_reports is not None
        ]

        self._db_session.execute(
            delete(CustomerIncomeBucket).where(
                CustomerIncomeBucket.customer_id.in_(customer_ids)
            )
        )
        if monthly_report_customer_ids:
            self._db_session.execute(
                delete(CustomerIncomeMonthlyReport).where(
                    CustomerIncomeMonthlyReport.customer_id.in_(monthly_report_customer_ids)
                )
            )

        for item in dto.items:
            income_summary = summary_map.get(item.customer_id)
            if income_summary is None:
                income_summary = CustomerIncomeSummary(customer_id=item.customer_id)
                self._db_session.add(income_summary)
                summary_map[item.customer_id] = income_summary

            income_summary.registered_income_amount = item.summary.registered_income_amount
            income_summary.issued_bill_count = item.summary.issued_bill_count
            income_summary.paid_bill_count = item.summary.paid_bill_count
            income_summary.collection_rate_percent = item.summary.collection_rate_percent

            for bucket in item.buckets:
                self._db_session.add(
                    CustomerIncomeBucket(
                        customer_id=item.customer_id,
                        bucket_code=bucket.bucket_code,
                        bucket_name=bucket.bucket_name,
                        registered_income_amount=bucket.registered_income_amount,
                    )
                )

            if item.monthly_reports is not None:
                for monthly_report in item.monthly_reports:
                    self._db_session.add(
                        CustomerIncomeMonthlyReport(
                            customer_id=item.customer_id,
                            month=monthly_report.month,
                            registered_income_amount=monthly_report.registered_income_amount,
                            issued_bill_count=monthly_report.issued_bill_count,
                            paid_bill_count=monthly_report.paid_bill_count,
                            collection_rate_percent=monthly_report.collection_rate_percent,
                        )
                    )

        self._commit_with_integrity_guard()

        for customer_id in customer_ids:
            self._db_session.refresh(summary_map[customer_id])

        bucket_map = self._get_bucket_map(customer_ids=customer_ids)
        monthly_report_map = self._get_monthly_report_map(customer_ids=customer_ids)
        return [
            (
                customer_map[item.customer_id],
                summary_map[item.customer_id],
                bucket_map.get(item.customer_id, []),
                monthly_report_map.get(item.customer_id, []),
            )
            for item in dto.items
        ]

    def _get_income_summary_map(
        self,
        *,
        customer_ids: list[int],
    ) -> dict[int, CustomerIncomeSummary]:
        summaries = list(
            self._db_session.scalars(
                select(CustomerIncomeSummary).where(
                    CustomerIncomeSummary.customer_id.in_(customer_ids)
                )
            ).all()
        )
        return {summary.customer_id: summary for summary in summaries}

    def _get_bucket_map(
        self,
        *,
        customer_ids: list[int],
    ) -> dict[int, list[CustomerIncomeBucket]]:
        buckets = list(
            self._db_session.scalars(
                select(CustomerIncomeBucket)
                .where(CustomerIncomeBucket.customer_id.in_(customer_ids))
                .order_by(
                    CustomerIncomeBucket.customer_id.asc(),
                    CustomerIncomeBucket.bucket_code.asc(),
                    CustomerIncomeBucket.id.asc(),
                )
            ).all()
        )
        bucket_map: dict[int, list[CustomerIncomeBucket]] = {}
        for bucket in buckets:
            bucket_map.setdefault(bucket.customer_id, []).append(bucket)
        return bucket_map

    def _get_monthly_report_map(
        self,
        *,
        customer_ids: list[int],
    ) -> dict[int, list[CustomerIncomeMonthlyReport]]:
        monthly_reports = list(
            self._db_session.scalars(
                select(CustomerIncomeMonthlyReport)
                .where(CustomerIncomeMonthlyReport.customer_id.in_(customer_ids))
                .order_by(
                    CustomerIncomeMonthlyReport.customer_id.asc(),
                    CustomerIncomeMonthlyReport.month.asc(),
                    CustomerIncomeMonthlyReport.id.asc(),
                )
            ).all()
        )
        monthly_report_map: dict[int, list[CustomerIncomeMonthlyReport]] = {}
        for monthly_report in monthly_reports:
            monthly_report_map.setdefault(monthly_report.customer_id, []).append(monthly_report)
        return monthly_report_map

    def _commit_with_integrity_guard(self) -> None:
        try:
            self._db_session.commit()
        except IntegrityError as exc:
            self._db_session.rollback()
            raise HTTPException(
                status_code=status.HTTP_409_CONFLICT,
                detail=DATA_INTEGRITY_ERROR,
            ) from exc


class CustomerIncomeImportService:
    SUMMARY_COLUMN_FALLBACKS = {
        "registered_income_amount": 0,
        "issued_bill_count": 1,
        "paid_bill_count": 2,
        "collection_rate_percent": 3,
    }
    SUMMARY_HEADER_ALIASES = {
        "registered_income_amount": (
            "registered_income_amount",
            "registered_income_amount_12m",
        ),
        "issued_bill_count": (
            "issued_bill_count",
            "issued_bills_count_12m",
        ),
        "paid_bill_count": (
            "paid_bill_count",
            "paid_bills_count_12m",
        ),
        "collection_rate_percent": (
            "collection_rate_percent",
            "collection_rate_percent_12m",
        ),
    }

    def __init__(self, db_session: Session) -> None:
        self._db_session = db_session

    def import_workbook(
        self,
        *,
        workbook_path: str | Path,
    ) -> CustomerIncomeImportReport:
        workbook_data = self.parse_workbook(workbook_path=workbook_path)
        summaries_by_name = {
            item.customer_name: item
            for item in workbook_data.summaries
        }
        buckets_by_name: dict[str, list[ImportedCustomerIncomeBucket]] = {}
        for bucket in workbook_data.buckets:
            buckets_by_name.setdefault(bucket.customer_name, []).append(bucket)

        customer_names = sorted(set(summaries_by_name) | set(buckets_by_name))
        upserted_summary_count = 0
        upserted_bucket_count = 0

        try:
            for customer_name in customer_names:
                customer = self._get_or_create_customer(customer_name=customer_name)
                summary_row = summaries_by_name.get(customer_name)
                income_summary = self._db_session.get(CustomerIncomeSummary, customer.id)
                if income_summary is None:
                    income_summary = CustomerIncomeSummary(customer_id=customer.id)
                    self._db_session.add(income_summary)
                income_summary.registered_income_amount = (
                    summary_row.registered_income_amount if summary_row else None
                )
                income_summary.issued_bill_count = (
                    summary_row.issued_bill_count if summary_row else None
                )
                income_summary.paid_bill_count = (
                    summary_row.paid_bill_count if summary_row else None
                )
                income_summary.collection_rate_percent = (
                    summary_row.collection_rate_percent if summary_row else None
                )
                upserted_summary_count += 1

                self._db_session.execute(
                    delete(CustomerIncomeBucket).where(
                        CustomerIncomeBucket.customer_id == customer.id
                    )
                )
                customer_buckets = buckets_by_name.get(customer_name, [])
                for bucket_row in customer_buckets:
                    self._db_session.add(
                        CustomerIncomeBucket(
                            customer_id=customer.id,
                            bucket_code=bucket_row.bucket_code,
                            bucket_name=bucket_row.bucket_name,
                            registered_income_amount=bucket_row.registered_income_amount,
                        )
                    )
                upserted_bucket_count += len(customer_buckets)
            self._db_session.commit()
        except Exception:
            self._db_session.rollback()
            raise

        return CustomerIncomeImportReport(
            processed_customer_count=len(customer_names),
            upserted_summary_count=upserted_summary_count,
            upserted_bucket_count=upserted_bucket_count,
            validation_failures=workbook_data.validation_failures,
        )

    def parse_workbook(
        self,
        *,
        workbook_path: str | Path,
    ) -> CustomerIncomeWorkbookData:
        workbook_file = Path(workbook_path)
        if not workbook_file.exists():
            raise FileNotFoundError(f"Workbook not found: {workbook_file}")

        workbook = load_workbook(filename=workbook_file, data_only=True)
        try:
            summary_sheet = self._get_sheet(workbook=workbook, sheet_name="Sheet1", fallback_index=0)
            bucket_sheet = self._get_sheet(workbook=workbook, sheet_name="Sheet2", fallback_index=1)
            summaries = self._parse_summary_sheet(sheet=summary_sheet)
            buckets = self._parse_bucket_sheet(sheet=bucket_sheet)
        finally:
            workbook.close()

        return CustomerIncomeWorkbookData(
            summaries=summaries,
            buckets=buckets,
            validation_failures=self._build_validation_failures(
                summaries=summaries,
                buckets=buckets,
            ),
        )

    def _get_or_create_customer(self, *, customer_name: str) -> Customer:
        customer = self._db_session.scalar(
            select(Customer)
            .where(Customer.name == customer_name)
            .order_by(Customer.id.asc())
            .limit(1)
        )
        if customer is not None:
            if not customer.is_active:
                customer.is_active = True
            return customer
        customer = Customer(
            name=customer_name,
            manager_name=None,
            grade=DEFAULT_IMPORTED_CUSTOMER_GRADE,
            is_active=True,
        )
        self._db_session.add(customer)
        self._db_session.flush()
        return customer

    def _get_sheet(self, *, workbook, sheet_name: str, fallback_index: int):
        if sheet_name in workbook.sheetnames:
            return workbook[sheet_name]
        if len(workbook.worksheets) <= fallback_index:
            raise ValueError(f"Worksheet '{sheet_name}' was not found in workbook.")
        return workbook.worksheets[fallback_index]

    def _parse_summary_sheet(self, *, sheet) -> list[ImportedCustomerIncomeSummary]:
        rows = list(sheet.iter_rows(values_only=True))
        if len(rows) <= 1:
            return []

        data_rows = rows[1:]
        municipality_name_index = self._detect_summary_name_column(data_rows=data_rows)
        header_lookup = {
            self._normalize_text(header).lower(): index
            for index, header in enumerate(rows[0])
            if self._normalize_text(header) is not None
        }
        parsed_rows: list[ImportedCustomerIncomeSummary] = []
        for row in data_rows:
            customer_name = self._normalize_text(
                row[municipality_name_index] if municipality_name_index < len(row) else None
            )
            if customer_name is None:
                continue
            parsed_rows.append(
                ImportedCustomerIncomeSummary(
                    customer_name=customer_name,
                    registered_income_amount=self._parse_int(
                        self._value_from_row(
                            row=row,
                            header_lookup=header_lookup,
                            field_name="registered_income_amount",
                        )
                    ),
                    issued_bill_count=self._parse_int(
                        self._value_from_row(
                            row=row,
                            header_lookup=header_lookup,
                            field_name="issued_bill_count",
                        )
                    ),
                    paid_bill_count=self._parse_int(
                        self._value_from_row(
                            row=row,
                            header_lookup=header_lookup,
                            field_name="paid_bill_count",
                        )
                    ),
                    collection_rate_percent=self._parse_float(
                        self._value_from_row(
                            row=row,
                            header_lookup=header_lookup,
                            field_name="collection_rate_percent",
                        )
                    ),
                )
            )
        return parsed_rows

    def _parse_bucket_sheet(self, *, sheet) -> list[ImportedCustomerIncomeBucket]:
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []

        current_customer_name = self._normalize_text(
            rows[0][3] if len(rows[0]) > 3 else None
        )
        raw_rows: list[ImportedCustomerIncomeBucket] = []
        for row in rows:
            if self._is_blank_row(row=row):
                current_customer_name = None
                continue
            row_customer_name = self._normalize_text(row[3] if len(row) > 3 else None)
            if row_customer_name is not None:
                current_customer_name = row_customer_name
            bucket_code = self._normalize_bucket_code(row[0] if len(row) > 0 else None)
            if bucket_code is None or bucket_code.lower() == "bucket_code":
                continue
            if current_customer_name is None:
                continue
            raw_rows.append(
                ImportedCustomerIncomeBucket(
                    customer_name=current_customer_name,
                    bucket_code=bucket_code,
                    bucket_name=self._normalize_text(row[1] if len(row) > 1 else None),
                    registered_income_amount=self._parse_int(row[2] if len(row) > 2 else None),
                )
            )

        canonical_bucket_names = self._build_canonical_bucket_names(bucket_rows=raw_rows)
        return [
            ImportedCustomerIncomeBucket(
                customer_name=item.customer_name,
                bucket_code=item.bucket_code,
                bucket_name=self._resolve_bucket_name(
                    raw_name=item.bucket_name,
                    canonical_name=canonical_bucket_names.get(item.bucket_code),
                ),
                registered_income_amount=item.registered_income_amount,
            )
            for item in raw_rows
        ]

    def _build_validation_failures(
        self,
        *,
        summaries: list[ImportedCustomerIncomeSummary],
        buckets: list[ImportedCustomerIncomeBucket],
    ) -> list[CustomerIncomeValidationFailure]:
        summaries_by_name = {
            item.customer_name: item
            for item in summaries
        }
        bucket_customer_names = {item.customer_name for item in buckets}
        bucket_totals: dict[str, int] = {}
        for bucket in buckets:
            if bucket.registered_income_amount is None:
                continue
            bucket_totals[bucket.customer_name] = (
                bucket_totals.get(bucket.customer_name, 0)
                + bucket.registered_income_amount
            )

        failures: list[CustomerIncomeValidationFailure] = []
        for customer_name in sorted(set(summaries_by_name) | bucket_customer_names):
            summary_row = summaries_by_name.get(customer_name)
            summary_amount = (
                summary_row.registered_income_amount
                if summary_row is not None
                else None
            )
            bucket_total = bucket_totals.get(customer_name)
            if summary_row is None and customer_name in bucket_customer_names:
                failures.append(
                    CustomerIncomeValidationFailure(
                        customer_name=customer_name,
                        message="Summary row missing in Sheet1.",
                        summary_registered_income_amount=None,
                        bucket_total_registered_income_amount=bucket_total,
                    )
                )
                continue
            if summary_row is not None and customer_name not in bucket_customer_names:
                failures.append(
                    CustomerIncomeValidationFailure(
                        customer_name=customer_name,
                        message="Bucket breakdown rows missing in Sheet2.",
                        summary_registered_income_amount=summary_amount,
                        bucket_total_registered_income_amount=None,
                    )
                )
                continue
            if summary_amount is not None and bucket_total is not None and summary_amount != bucket_total:
                failures.append(
                    CustomerIncomeValidationFailure(
                        customer_name=customer_name,
                        message="Bucket total does not match summary registered income.",
                        summary_registered_income_amount=summary_amount,
                        bucket_total_registered_income_amount=bucket_total,
                    )
                )
        return failures

    def _detect_summary_name_column(self, *, data_rows: list[tuple[object, ...]]) -> int:
        candidate_index = -1
        max_length = max((len(row) for row in data_rows), default=0)
        for index in range(max_length):
            has_text_value = any(
                self._is_textual_value(row[index] if index < len(row) else None)
                for row in data_rows
            )
            if has_text_value:
                candidate_index = index
        if candidate_index >= 0:
            return candidate_index
        return 4

    def _value_from_row(
        self,
        *,
        row: tuple[object, ...],
        header_lookup: dict[str, int],
        field_name: str,
    ) -> object | None:
        column_index = self._resolve_summary_column_index(
            header_lookup=header_lookup,
            field_name=field_name,
        )
        if column_index >= len(row):
            return None
        return row[column_index]

    def _resolve_summary_column_index(
        self,
        *,
        header_lookup: dict[str, int],
        field_name: str,
    ) -> int:
        for header_name in self.SUMMARY_HEADER_ALIASES[field_name]:
            column_index = header_lookup.get(header_name)
            if column_index is not None:
                return column_index
        return self.SUMMARY_COLUMN_FALLBACKS[field_name]

    def _build_canonical_bucket_names(
        self,
        *,
        bucket_rows: list[ImportedCustomerIncomeBucket],
    ) -> dict[str, str]:
        canonical_bucket_names: dict[str, str] = {}
        for row in bucket_rows:
            if row.bucket_name is None or self._is_numeric_like(row.bucket_name):
                continue
            canonical_bucket_names.setdefault(row.bucket_code, row.bucket_name)
        return canonical_bucket_names

    def _resolve_bucket_name(
        self,
        *,
        raw_name: str | None,
        canonical_name: str | None,
    ) -> str | None:
        if raw_name is None or self._is_numeric_like(raw_name):
            return canonical_name or raw_name
        return raw_name

    def _normalize_bucket_code(self, value: object | None) -> str | None:
        normalized_text = self._normalize_text(value)
        if normalized_text is None:
            return None
        if self._is_numeric_like(normalized_text):
            parsed_value = self._parse_int(normalized_text)
            if parsed_value is not None:
                return str(parsed_value)
        return normalized_text.upper()

    def _normalize_text(self, value: object | None) -> str | None:
        if value is None:
            return None
        text_value = str(value).strip()
        if not text_value:
            return None
        return text_value

    def _parse_int(self, value: object | None) -> int | None:
        if value is None:
            return None
        if isinstance(value, bool):
            return int(value)
        if isinstance(value, int):
            return value
        if isinstance(value, float):
            return int(value)
        try:
            return int(Decimal(str(value).replace(",", "").strip()))
        except (InvalidOperation, ValueError):
            return None

    def _parse_float(self, value: object | None) -> float | None:
        if value is None:
            return None
        if isinstance(value, bool):
            return float(value)
        if isinstance(value, (int, float)):
            return float(value)
        try:
            return float(Decimal(str(value).replace(",", "").strip()))
        except (InvalidOperation, ValueError):
            return None

    def _is_blank_row(self, *, row: tuple[object, ...]) -> bool:
        return all(self._normalize_text(value) is None for value in row)

    def _is_textual_value(self, value: object | None) -> bool:
        normalized_text = self._normalize_text(value)
        if normalized_text is None:
            return False
        return not self._is_numeric_like(normalized_text)

    def _is_numeric_like(self, value: object | None) -> bool:
        normalized_text = self._normalize_text(value)
        if normalized_text is None:
            return False
        try:
            Decimal(normalized_text.replace(",", ""))
        except InvalidOperation:
            return False
        return True


@dataclass(frozen=True)
class ResolvedCustomerBridgeConfig:
    instance_id: str
    base_url_internal: str
    audience: str
    tenant_id: str
    timeout_seconds: int
    retry_count: int
    retry_backoff_seconds: float
    legacy_api_key: str | None = None


class CustomerBridgeConfigResolver:
    def __init__(self, settings: Settings) -> None:
        self._settings = settings

    def resolve(
        self,
        customer: Customer,
        bridge_config: CustomerBridgeConfig | None,
    ) -> ResolvedCustomerBridgeConfig:
        missing_fields: list[str] = []
        if bridge_config is None:
            missing_fields.append("bridge_row")
        else:
            if not self._is_active_status(bridge_config.status):
                missing_fields.append("status")
            if not bridge_config.instance_id:
                missing_fields.append("instance_id")
            if not bridge_config.base_url_internal:
                missing_fields.append("base_url_internal")
            if not bridge_config.audience:
                missing_fields.append("audience")
            if not bridge_config.tenant_id:
                missing_fields.append("tenant_id")
        if missing_fields:
            raise HTTPException(
                status_code=status.HTTP_409_CONFLICT,
                detail={
                    "message": CUSTOMER_BRIDGE_NOT_CONFIGURED,
                    "developer_message": (
                        f"Customer {customer.id} bridge configuration is incomplete. "
                        f"Missing or disabled fields: {', '.join(missing_fields)}."
                    ),
                },
            )
        timeout_seconds = (
            bridge_config.request_timeout_seconds
            if bridge_config.request_timeout_seconds is not None
            else self._settings.bridge_request_timeout_seconds
        )
        retry_count = (
            bridge_config.request_retry_count
            if bridge_config.request_retry_count is not None
            else self._settings.bridge_request_retry_count
        )
        retry_backoff_seconds = (
            bridge_config.request_retry_backoff_seconds
            if bridge_config.request_retry_backoff_seconds is not None
            else self._settings.bridge_request_retry_backoff_seconds
        )
        return ResolvedCustomerBridgeConfig(
            instance_id=bridge_config.instance_id,
            base_url_internal=bridge_config.base_url_internal,
            audience=bridge_config.audience,
            tenant_id=bridge_config.tenant_id,
            timeout_seconds=timeout_seconds,
            retry_count=retry_count,
            retry_backoff_seconds=retry_backoff_seconds,
            legacy_api_key=bridge_config.bridge_api_key,
        )

    def _is_active_status(self, status_value: str | None) -> bool:
        if status_value is None:
            return False
        return status_value.strip().lower() == "active"


class CustomerBridgeService:
    def __init__(
        self,
        bridge_config_service: CustomerBridgeConfigService,
        bridge_client: BridgeClient,
        bridge_access_token_service: BridgeAccessTokenService,
        settings: Settings,
    ) -> None:
        self._bridge_config_service = bridge_config_service
        self._bridge_client = bridge_client
        self._bridge_access_token_service = bridge_access_token_service
        self._config_resolver = CustomerBridgeConfigResolver(settings=settings)

    def get_health(
        self,
        customer_id: int,
        instance_id: str | None = None,
        correlation_id: str | None = None,
    ) -> tuple[Customer, CustomerBridgeConfig | None, BridgeHealthResult]:
        customer, bridge_config, request = self._build_request(
            customer_id=customer_id,
            instance_id=instance_id,
            correlation_id=correlation_id,
        )
        checked_at = datetime.now(timezone.utc)
        try:
            bridge_health = self._bridge_client.get_health(request=request)
        except (
            BridgeConnectionError,
            BridgeUnauthorizedError,
            BridgeUnexpectedStatusError,
            BridgeInvalidResponseError,
        ) as exc:
            if bridge_config is not None:
                bridge_config = self._bridge_config_service.persist_health_status(
                    bridge_config=bridge_config,
                    is_online=False,
                    checked_at=checked_at,
                    error=self._build_health_error_message(exc),
                )
            raise self._map_bridge_error(
                customer=customer,
                bridge_base_url=bridge_config.base_url_internal if bridge_config else None,
                exc=exc,
            ) from exc
        if bridge_config is not None:
            bridge_config = self._bridge_config_service.persist_health_status(
                bridge_config=bridge_config,
                is_online=True,
                checked_at=checked_at,
                error=None,
            )
        return customer, bridge_config, bridge_health

    def refresh_status(
        self,
        customer_id: int,
        instance_id: str | None = None,
        correlation_id: str | None = None,
    ) -> tuple[Customer, CustomerBridgeConfig | None]:
        customer, bridge_config, request = self._build_request(
            customer_id=customer_id,
            instance_id=instance_id,
            correlation_id=correlation_id,
        )
        checked_at = datetime.now(timezone.utc)
        try:
            self._bridge_client.get_health(request=request)
        except (
            BridgeConnectionError,
            BridgeUnauthorizedError,
            BridgeUnexpectedStatusError,
            BridgeInvalidResponseError,
        ) as exc:
            if bridge_config is not None:
                bridge_config = self._bridge_config_service.persist_health_status(
                    bridge_config=bridge_config,
                    is_online=False,
                    checked_at=checked_at,
                    error=self._build_health_error_message(exc),
                )
            return customer, bridge_config
        if bridge_config is not None:
            bridge_config = self._bridge_config_service.persist_health_status(
                bridge_config=bridge_config,
                is_online=True,
                checked_at=checked_at,
                error=None,
            )
        return customer, bridge_config

    def get_capabilities(
        self,
        customer_id: int,
        instance_id: str | None = None,
        correlation_id: str | None = None,
    ) -> tuple[Customer, CustomerBridgeConfig | None, BridgeCapabilitiesResult]:
        customer, bridge_config, request = self._build_request(
            customer_id=customer_id,
            instance_id=instance_id,
            correlation_id=correlation_id,
        )
        try:
            bridge_capabilities = self._bridge_client.get_capabilities(request=request)
        except (
            BridgeConnectionError,
            BridgeUnauthorizedError,
            BridgeUnexpectedStatusError,
            BridgeInvalidResponseError,
        ) as exc:
            raise self._map_bridge_error(
                customer=customer,
                bridge_base_url=bridge_config.base_url_internal if bridge_config else None,
                exc=exc,
            ) from exc
        return customer, bridge_config, bridge_capabilities

    def fetch_active_subscription(
        self,
        customer_id: int,
        instance_id: str | None = None,
        correlation_id: str | None = None,
    ) -> tuple[Customer, CustomerBridgeConfig | None, BridgeSubscriptionResult]:
        customer, bridge_config, request = self._build_request(
            customer_id=customer_id,
            instance_id=instance_id,
            correlation_id=correlation_id,
        )
        checked_at = datetime.now(timezone.utc)
        try:
            bridge_subscription = self._bridge_client.get_active_subscription(request=request)
        except (
            BridgeConnectionError,
            BridgeUnauthorizedError,
            BridgeUnexpectedStatusError,
            BridgeInvalidResponseError,
        ) as exc:
            self._handle_subscription_bridge_error(
                bridge_config=bridge_config,
                checked_at=checked_at,
                exc=exc,
            )
            raise self._map_subscription_error(
                customer=customer,
                bridge_base_url=bridge_config.base_url_internal if bridge_config else None,
                exc=exc,
            ) from exc
        if bridge_config is not None:
            bridge_config = self._bridge_config_service.persist_subscription_snapshot(
                bridge_config=bridge_config,
                subscription=bridge_subscription,
                checked_at=checked_at,
            )
        return customer, bridge_config, bridge_subscription

    def update_bridge_subscription(
        self,
        customer_id: int,
        payload: dict[str, object],
        instance_id: str | None = None,
        correlation_id: str | None = None,
    ) -> tuple[Customer, CustomerBridgeConfig | None, BridgeSubscriptionResult]:
        customer, bridge_config, request = self._build_request(
            customer_id=customer_id,
            instance_id=instance_id,
            correlation_id=correlation_id,
        )
        checked_at = datetime.now(timezone.utc)
        try:
            bridge_subscription = self._bridge_client.update_active_subscription(
                request=request,
                payload=payload,
            )
        except (
            BridgeConnectionError,
            BridgeUnauthorizedError,
            BridgeUnexpectedStatusError,
            BridgeInvalidResponseError,
        ) as exc:
            self._handle_subscription_bridge_error(
                bridge_config=bridge_config,
                checked_at=checked_at,
                exc=exc,
            )
            raise self._map_subscription_error(
                customer=customer,
                bridge_base_url=bridge_config.base_url_internal if bridge_config else None,
                exc=exc,
            ) from exc
        if bridge_config is not None:
            bridge_config = self._bridge_config_service.persist_subscription_snapshot(
                bridge_config=bridge_config,
                subscription=bridge_subscription,
                checked_at=checked_at,
            )
        return customer, bridge_config, bridge_subscription

    def get_subscription_messages(
        self,
        customer_id: int,
        instance_id: str | None = None,
        correlation_id: str | None = None,
    ) -> tuple[Customer, CustomerBridgeConfig | None, list[BridgeSubscriptionMessageResult]]:
        customer, bridge_config, request = self._build_request(
            customer_id=customer_id,
            instance_id=instance_id,
            correlation_id=correlation_id,
        )
        try:
            bridge_messages = self._bridge_client.get_subscription_messages(request=request)
        except (
            BridgeConnectionError,
            BridgeUnauthorizedError,
            BridgeUnexpectedStatusError,
            BridgeInvalidResponseError,
        ) as exc:
            raise self._map_bridge_error(
                customer=customer,
                bridge_base_url=bridge_config.base_url_internal if bridge_config else None,
                exc=exc,
            ) from exc
        return customer, bridge_config, bridge_messages

    def upsert_subscription_messages(
        self,
        customer_id: int,
        payload: list[dict[str, object]],
        instance_id: str | None = None,
        correlation_id: str | None = None,
    ) -> tuple[Customer, CustomerBridgeConfig | None, list[BridgeSubscriptionMessageResult]]:
        customer, bridge_config, request = self._build_request(
            customer_id=customer_id,
            instance_id=instance_id,
            correlation_id=correlation_id,
        )
        try:
            bridge_messages = self._bridge_client.upsert_subscription_messages(
                request=request,
                payload=payload,
            )
        except (
            BridgeConnectionError,
            BridgeUnauthorizedError,
            BridgeUnexpectedStatusError,
            BridgeInvalidResponseError,
        ) as exc:
            raise self._map_bridge_error(
                customer=customer,
                bridge_base_url=bridge_config.base_url_internal if bridge_config else None,
                exc=exc,
            ) from exc
        return customer, bridge_config, bridge_messages

    def get_subscription_config(
        self,
        customer_id: int,
        instance_id: str | None = None,
        correlation_id: str | None = None,
    ) -> tuple[Customer, CustomerBridgeConfig | None, BridgeSubscriptionConfigResult]:
        customer, bridge_config, request = self._build_request(
            customer_id=customer_id,
            instance_id=instance_id,
            correlation_id=correlation_id,
        )
        checked_at = datetime.now(timezone.utc)
        try:
            bridge_config_result = self._bridge_client.get_subscription_config(request=request)
        except (
            BridgeConnectionError,
            BridgeUnauthorizedError,
            BridgeUnexpectedStatusError,
            BridgeInvalidResponseError,
        ) as exc:
            self._handle_subscription_bridge_error(
                bridge_config=bridge_config,
                checked_at=checked_at,
                exc=exc,
            )
            raise self._map_subscription_error(
                customer=customer,
                bridge_base_url=bridge_config.base_url_internal if bridge_config else None,
                exc=exc,
            ) from exc
        if bridge_config is not None:
            bridge_config = self._bridge_config_service.persist_subscription_snapshot(
                bridge_config=bridge_config,
                subscription=bridge_config_result.subscription,
                checked_at=checked_at,
            )
        return customer, bridge_config, bridge_config_result

    def sync_subscription_config(
        self,
        customer_id: int,
        payload: dict[str, object],
        instance_id: str | None = None,
        correlation_id: str | None = None,
    ) -> tuple[Customer, CustomerBridgeConfig | None, BridgeSubscriptionConfigResult]:
        customer, bridge_config, request = self._build_request(
            customer_id=customer_id,
            instance_id=instance_id,
            correlation_id=correlation_id,
        )
        checked_at = datetime.now(timezone.utc)
        try:
            bridge_config_result = self._bridge_client.sync_subscription_config(
                request=request,
                payload=payload,
            )
        except (
            BridgeConnectionError,
            BridgeUnauthorizedError,
            BridgeUnexpectedStatusError,
            BridgeInvalidResponseError,
        ) as exc:
            self._handle_subscription_bridge_error(
                bridge_config=bridge_config,
                checked_at=checked_at,
                exc=exc,
            )
            raise self._map_subscription_error(
                customer=customer,
                bridge_base_url=bridge_config.base_url_internal if bridge_config else None,
                exc=exc,
            ) from exc
        if bridge_config is not None:
            bridge_config = self._bridge_config_service.persist_subscription_snapshot(
                bridge_config=bridge_config,
                subscription=bridge_config_result.subscription,
                checked_at=checked_at,
            )
        return customer, bridge_config, bridge_config_result

    def get_subscription(
        self,
        customer_id: int,
        instance_id: str | None = None,
    ) -> tuple[Customer, CustomerBridgeConfig]:
        customer, bridge_config = self._bridge_config_service.get_active(
            customer_id=customer_id,
            instance_id=instance_id,
        )
        self._config_resolver.resolve(
            customer=customer,
            bridge_config=bridge_config,
        )
        if bridge_config is None:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail={
                    "message": CUSTOMER_BRIDGE_SUBSCRIPTION_NOT_CACHED,
                    "developer_message": (
                        f"Customer {customer.id} does not have a bridge configuration record."
                    ),
                },
            )
        has_cached_subscription = (
            bridge_config.cached_subscription_start_date is not None
            and bridge_config.cached_subscription_end_date is not None
            and bridge_config.cached_subscription_is_active is not None
        )
        if has_cached_subscription:
            return customer, bridge_config
        cached_message = (
            bridge_config.last_subscription_error
            or CUSTOMER_BRIDGE_SUBSCRIPTION_NOT_CACHED
        )
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail={
                "message": cached_message,
                "developer_message": (
                    f"Customer {customer.id} does not have a cached bridge subscription. "
                    "Run the subscription refresh endpoint first."
                ),
            },
        )

    def refresh_subscription(
        self,
        customer_id: int,
        instance_id: str | None = None,
        correlation_id: str | None = None,
    ) -> tuple[Customer, CustomerBridgeConfig | None]:
        customer, bridge_config, request = self._build_request(
            customer_id=customer_id,
            instance_id=instance_id,
            correlation_id=correlation_id,
        )
        checked_at = datetime.now(timezone.utc)
        try:
            bridge_subscription = self._bridge_client.get_active_subscription(
                request=request,
            )
        except (
            BridgeConnectionError,
            BridgeUnauthorizedError,
            BridgeUnexpectedStatusError,
            BridgeInvalidResponseError,
        ) as exc:
            if bridge_config is not None:
                if (
                    isinstance(exc, BridgeUnexpectedStatusError)
                    and exc.status_code == status.HTTP_404_NOT_FOUND
                ):
                    bridge_config = self._bridge_config_service.persist_subscription_absence(
                        bridge_config=bridge_config,
                        checked_at=checked_at,
                        error=self._build_subscription_error_message(exc),
                    )
                else:
                    bridge_config = self._bridge_config_service.persist_subscription_refresh_error(
                        bridge_config=bridge_config,
                        error=self._build_subscription_error_message(exc),
                    )
            raise self._map_subscription_error(
                customer=customer,
                bridge_base_url=bridge_config.base_url_internal if bridge_config else None,
                exc=exc,
            ) from exc
        if bridge_config is not None:
            bridge_config = self._bridge_config_service.persist_subscription_snapshot(
                bridge_config=bridge_config,
                subscription=bridge_subscription,
                checked_at=checked_at,
            )
        return customer, bridge_config

    def _build_request(
        self,
        customer_id: int,
        instance_id: str | None,
        correlation_id: str | None,
    ) -> tuple[Customer, CustomerBridgeConfig | None, BridgeRequest]:
        customer, bridge_config = self._bridge_config_service.get_active(
            customer_id=customer_id,
            instance_id=instance_id,
        )
        resolved_bridge_config = self._config_resolver.resolve(
            customer=customer,
            bridge_config=bridge_config,
        )
        bridge_access_token = self._bridge_access_token_service.create_access_token(
            customer_id=customer.id,
            instance_id=resolved_bridge_config.instance_id,
            tenant_id=resolved_bridge_config.tenant_id,
            audience=resolved_bridge_config.audience,
        )
        return customer, bridge_config, BridgeRequest(
            base_url=resolved_bridge_config.base_url_internal,
            access_token=bridge_access_token,
            timeout_seconds=resolved_bridge_config.timeout_seconds,
            correlation_id=correlation_id,
            retry_count=resolved_bridge_config.retry_count,
            retry_backoff_seconds=resolved_bridge_config.retry_backoff_seconds,
            legacy_api_key=resolved_bridge_config.legacy_api_key,
        )

    def _map_bridge_error(
        self,
        customer: Customer,
        bridge_base_url: str | None,
        exc: (
            BridgeConnectionError
            | BridgeUnauthorizedError
            | BridgeUnexpectedStatusError
            | BridgeInvalidResponseError
        ),
    ) -> HTTPException:
        if isinstance(exc, BridgeConnectionError):
            return HTTPException(
                status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
                detail={
                    "message": CUSTOMER_BRIDGE_UNAVAILABLE,
                    "developer_message": (
                        f"Bridge request for customer {customer.id} could not reach "
                        f"{bridge_base_url}: {exc}"
                    ),
                },
            )
        if isinstance(exc, BridgeUnauthorizedError):
            debug_context = self._build_upstream_debug_context(
                status_code=exc.status_code,
                response_body=exc.response_body,
            )
            return HTTPException(
                status_code=status.HTTP_502_BAD_GATEWAY,
                detail={
                    "message": CUSTOMER_BRIDGE_AUTH_FAILED,
                    "developer_message": (
                        f"Bridge request for customer {customer.id} was rejected with "
                        f"status {exc.status_code} at {bridge_base_url}. "
                        f"{debug_context}"
                    ),
                },
            )
        if isinstance(exc, BridgeUnexpectedStatusError):
            debug_context = self._build_upstream_debug_context(
                status_code=exc.status_code,
                response_body=exc.response_body,
            )
            return HTTPException(
                status_code=status.HTTP_502_BAD_GATEWAY,
                detail={
                    "message": CUSTOMER_BRIDGE_REQUEST_FAILED,
                    "developer_message": (
                        f"Bridge request for customer {customer.id} returned "
                        f"unexpected status {exc.status_code} at {bridge_base_url}. "
                        f"{debug_context}"
                    ),
                },
            )
        return HTTPException(
            status_code=status.HTTP_502_BAD_GATEWAY,
            detail={
                "message": CUSTOMER_BRIDGE_INVALID_RESPONSE,
                "developer_message": (
                    f"Bridge response for customer {customer.id} did not match the "
                    f"expected schema at {bridge_base_url}: {exc}"
                ),
            },
        )

    def refresh_all_statuses(self) -> list[tuple[Customer, CustomerBridgeConfig | None]]:
        results: list[tuple[Customer, CustomerBridgeConfig | None]] = []
        for customer_id in self._bridge_config_service.list_refreshable_customer_ids():
            results.append(self.refresh_status(customer_id=customer_id))
        return results

    def _map_subscription_error(
        self,
        customer: Customer,
        bridge_base_url: str | None,
        exc: (
            BridgeConnectionError
            | BridgeUnauthorizedError
            | BridgeUnexpectedStatusError
            | BridgeInvalidResponseError
        ),
    ) -> HTTPException:
        if (
            isinstance(exc, BridgeUnexpectedStatusError)
            and exc.status_code == status.HTTP_404_NOT_FOUND
        ):
            upstream_message = (
                self._extract_bridge_error_message(response_body=exc.response_body)
                or "هیچ اشتراک فعالی وجود ندارد."
            )
            debug_context = self._build_upstream_debug_context(
                status_code=exc.status_code,
                response_body=exc.response_body,
            )
            return HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail={
                    "message": upstream_message,
                    "developer_message": (
                        f"Bridge subscription request for customer {customer.id} "
                        f"at {bridge_base_url} returned status 404. "
                        f"{debug_context}"
                    ),
                },
            )
        return self._map_bridge_error(
            customer=customer,
            bridge_base_url=bridge_base_url,
            exc=exc,
        )

    def _build_health_error_message(
        self,
        exc: (
            BridgeConnectionError
            | BridgeUnauthorizedError
            | BridgeUnexpectedStatusError
            | BridgeInvalidResponseError
        ),
    ) -> str:
        return str(exc)

    def _build_subscription_error_message(
        self,
        exc: (
            BridgeConnectionError
            | BridgeUnauthorizedError
            | BridgeUnexpectedStatusError
            | BridgeInvalidResponseError
        ),
    ) -> str:
        if isinstance(exc, BridgeConnectionError):
            return CUSTOMER_BRIDGE_UNAVAILABLE
        if isinstance(exc, BridgeUnauthorizedError):
            return CUSTOMER_BRIDGE_AUTH_FAILED
        if isinstance(exc, BridgeUnexpectedStatusError):
            if exc.status_code == status.HTTP_404_NOT_FOUND:
                return (
                    self._extract_bridge_error_message(response_body=exc.response_body)
                    or "هیچ اشتراک فعالی وجود ندارد."
                )
            return CUSTOMER_BRIDGE_REQUEST_FAILED
        return CUSTOMER_BRIDGE_INVALID_RESPONSE

    def _handle_subscription_bridge_error(
        self,
        bridge_config: CustomerBridgeConfig | None,
        checked_at: datetime,
        exc: (
            BridgeConnectionError
            | BridgeUnauthorizedError
            | BridgeUnexpectedStatusError
            | BridgeInvalidResponseError
        ),
    ) -> None:
        if bridge_config is None:
            return
        if isinstance(exc, BridgeUnexpectedStatusError) and exc.status_code == status.HTTP_404_NOT_FOUND:
            self._bridge_config_service.persist_subscription_absence(
                bridge_config=bridge_config,
                checked_at=checked_at,
                error=self._build_subscription_error_message(exc),
            )
            return
        self._bridge_config_service.persist_subscription_refresh_error(
            bridge_config=bridge_config,
            error=self._build_subscription_error_message(exc),
        )

    def _extract_bridge_error_message(self, response_body: str | None) -> str | None:
        if not response_body:
            return None
        try:
            payload = json.loads(response_body)
        except json.JSONDecodeError:
            return None
        if not isinstance(payload, dict):
            return None
        for field_name in ("error", "message", "detail"):
            field_value = payload.get(field_name)
            if isinstance(field_value, str) and field_value.strip():
                return field_value.strip()
        return None

    def _build_upstream_debug_context(
        self,
        *,
        status_code: int,
        response_body: str | None,
    ) -> str:
        upstream_message = self._extract_bridge_error_message(response_body=response_body)
        body_excerpt = self._compact_response_body(response_body=response_body)
        if upstream_message and body_excerpt:
            if body_excerpt == upstream_message:
                return f"Upstream message: {upstream_message}"
            return (
                f"Upstream message: {upstream_message}. "
                f"Upstream response body excerpt: {body_excerpt}"
            )
        if upstream_message:
            return f"Upstream message: {upstream_message}"
        if body_excerpt:
            return f"Upstream response body excerpt: {body_excerpt}"
        return f"Upstream response body was empty for status {status_code}."

    def _compact_response_body(self, response_body: str | None) -> str | None:
        if not response_body:
            return None
        compact = " ".join(response_body.split())
        if not compact:
            return None
        max_length = 500
        if len(compact) <= max_length:
            return compact
        return f"{compact[:max_length]}..."


def get_customer_service(
    db_session: Session = Depends(get_db_session),
) -> CustomerService:
    return CustomerService(db_session=db_session)


def get_customer_income_service(
    db_session: Session = Depends(get_db_session),
    customer_service: CustomerService = Depends(get_customer_service),
) -> CustomerIncomeService:
    return CustomerIncomeService(
        db_session=db_session,
        customer_service=customer_service,
    )


def get_customer_bridge_config_service(
    db_session: Session = Depends(get_db_session),
    customer_service: CustomerService = Depends(get_customer_service),
) -> CustomerBridgeConfigService:
    return CustomerBridgeConfigService(
        db_session=db_session,
        customer_service=customer_service,
    )


def get_customer_bridge_service(
    customer_bridge_config_service: CustomerBridgeConfigService = Depends(get_customer_bridge_config_service),
    bridge_client: BridgeClient = Depends(get_bridge_client),
    bridge_access_token_service: BridgeAccessTokenService = Depends(get_bridge_access_token_service),
    settings: Settings = Depends(get_settings),
) -> CustomerBridgeService:
    return CustomerBridgeService(
        bridge_config_service=customer_bridge_config_service,
        bridge_client=bridge_client,
        bridge_access_token_service=bridge_access_token_service,
        settings=settings,
    )


def get_customer_database_connection_verifier() -> CustomerDatabaseConnectionVerifier:
    return CustomerDatabaseConnectionVerifier()


def get_customer_database_connection_service(
    db_session: Session = Depends(get_db_session),
    customer_service: CustomerService = Depends(get_customer_service),
    secret_cipher_service: SecretCipherService = Depends(get_secret_cipher_service),
    verifier: CustomerDatabaseConnectionVerifier = Depends(get_customer_database_connection_verifier),
) -> CustomerDatabaseConnectionService:
    return CustomerDatabaseConnectionService(
        db_session=db_session,
        customer_service=customer_service,
        secret_cipher_service=secret_cipher_service,
        verifier=verifier,
    )

