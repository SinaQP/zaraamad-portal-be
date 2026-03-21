from contextlib import contextmanager
from pathlib import Path

from openpyxl import Workbook, load_workbook

from app.commands.export_db_results import (
    DatabaseResultExporter,
    ExportSummary,
    ServerInputReader,
    SqlServerConnector,
    configure_logging,
    load_query_text,
    validate_read_only_query,
)


class FakeSqlServerConnector:
    def discover_databases(self, *, server):
        if server.host == "10.0.0.2":
            raise RuntimeError("server unavailable")
        return ["main_db", "empty_db"]

    @contextmanager
    def stream_query_rows(self, *, server, database_name: str, query: str):
        del server, query
        if database_name == "main_db":
            yield ["id", "server_ip"], iter([(1, "from_query")])
            return
        yield ["id"], iter([])


def test_server_input_reader_reads_xlsx_with_header_normalization(
    tmp_path: Path,
) -> None:
    workbook_path = tmp_path / "servers.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Host", "Server Port", "User Name", "Pwd"])
    sheet.append(["10.0.0.1", "1433", "sa", "secret"])
    sheet.append([None, "1433", "missing-host", "secret"])
    workbook.save(workbook_path)
    workbook.close()

    reader = ServerInputReader()
    servers, errors = reader.read(workbook_path)

    assert len(servers) == 1
    assert servers[0].host == "10.0.0.1"
    assert servers[0].port == 1433
    assert servers[0].username == "sa"
    assert len(errors) == 1
    assert "Missing ip/host value." in errors[0].error_message


def test_server_input_reader_reads_csv(tmp_path: Path) -> None:
    csv_path = tmp_path / "servers.csv"
    csv_path.write_text(
        "ip,port,username,password\n10.0.0.3,1444,report,pass123\n",
        encoding="utf-8",
    )

    reader = ServerInputReader()
    servers, errors = reader.read(csv_path)

    assert errors == []
    assert len(servers) == 1
    assert servers[0].host == "10.0.0.3"
    assert servers[0].port == 1444
    assert servers[0].username == "report"


def test_exporter_writes_results_and_errors_workbook(tmp_path: Path) -> None:
    input_path = tmp_path / "servers.csv"
    output_path = tmp_path / "db_results.xlsx"
    input_path.write_text(
        (
            "ip,port,username,password\n"
            "10.0.0.1,1433,sa,secret\n"
            "10.0.0.2,1433,sa,secret\n"
        ),
        encoding="utf-8",
    )

    exporter = DatabaseResultExporter(
        input_reader=ServerInputReader(),
        connector=FakeSqlServerConnector(),
        logger=configure_logging(log_level="INFO"),
    )
    summary = exporter.run(
        input_path=input_path,
        query="SELECT 1",
        output_path=output_path,
    )

    assert summary == ExportSummary(
        servers_processed=2,
        databases_scanned=2,
        success_count=2,
        error_count=1,
        output_path=output_path,
    )

    workbook = load_workbook(output_path)
    try:
        results_rows = list(workbook["results"].iter_rows(values_only=True))
        errors_rows = list(workbook["errors"].iter_rows(values_only=True))
    finally:
        workbook.close()

    assert results_rows[0] == (
        "server_ip",
        "port",
        "database_name",
        "scan_status",
        "id",
        "query_server_ip",
    )
    assert results_rows[1] == ("10.0.0.1", 1433, "main_db", "ok", 1, "from_query")
    assert results_rows[2] == ("10.0.0.1", 1433, "empty_db", "no_rows", None, None)

    assert errors_rows[0] == ("server_ip", "port", "database_name", "error_message")
    assert errors_rows[1][0:3] == ("10.0.0.2", 1433, None)
    assert "server unavailable" in errors_rows[1][3]


def test_load_query_text_supports_query_file(tmp_path: Path) -> None:
    query_file = tmp_path / "query.sql"
    query_file.write_text("SELECT TOP 1 name FROM sys.databases;", encoding="utf-8")

    assert load_query_text(inline_query=None, query_file=str(query_file)) == (
        "SELECT TOP 1 name FROM sys.databases;"
    )


def test_validate_read_only_query_allows_declare_and_select() -> None:
    query_text = """
    /* safe report query */
    DECLARE @start_date date = '2026-01-01';
    ;WITH sample AS (
        SELECT 1 AS id
    )
    SELECT id
    FROM sample
    WHERE id = 1;
    """

    assert validate_read_only_query(query_text) == query_text


def test_validate_read_only_query_rejects_write_statements() -> None:
    query_text = """
    SELECT 1;
    DELETE FROM dbo.Customers;
    """

    try:
        validate_read_only_query(query_text)
    except ValueError as exc:
        assert "delete" in str(exc).lower()
    else:
        raise AssertionError("Expected read-only validation to reject DELETE statements.")
